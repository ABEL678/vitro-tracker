# vitro/ui/categories.py
"""
🏷 Категории — аналитика + работа с замечаниями.

Под-вкладки:
  📈 Анализ — распределение, по дисциплинам, по авторам.
  📝 Работа с замечаниями — st.data_editor с каскадными фильтрами.
  📥 Импорт из Excel — массовая загрузка категорий из файлов.

Источник — `_load_all_categorized` (deadlines.py).
Логика: holder == "ours" — только замечания на нашей стороне.
"""

import io
from datetime import datetime, timedelta

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

from vitro.sqlite_db import (
    get_conn,
    load_remarks_for_editor,
    update_category_safe,
    update_category_force,
)
from vitro.disciplines import discipline_name


# ---------------------------------------------------------------------------
#  Константы
# ---------------------------------------------------------------------------
CAT_1 = "Принято/корректное"
CAT_2 = "Формальное/нет влияния на СМР"
CAT_3 = "Доп.требование/отсутствует в ТЗ"
CAT_4 = "Не принято/нарушение ТНПА"

CAT_PREFIX = {
    CAT_1: "🟢",
    CAT_2: "🟡",
    CAT_3: "🔵",
    CAT_4: "🔴",
}

CAT_OPTIONS_DISPLAY = [f"{CAT_PREFIX[c]} {c}" for c in [CAT_1, CAT_2, CAT_3, CAT_4]]
CAT_OPTIONS_RAW = [CAT_1, CAT_2, CAT_3, CAT_4]

CAT_SHORT = {
    CAT_1: "Принято",
    CAT_2: "Формальное",
    CAT_3: "Доп.треб.",
    CAT_4: "Не принято",
    "Без категории": "Без категории",
}

CAT_COLORS = {
    CAT_1:           "#C6EFCE",
    CAT_2:           "#FFEB9C",
    CAT_3:           "#BDD7EE",
    CAT_4:           "#FFC7CE",
    "Без категории": "#D9D9D9",
    "Принято":       "#C6EFCE",
    "Формальное":    "#FFEB9C",
    "Доп.треб.":     "#BDD7EE",
    "Не принято":    "#FFC7CE",
}

STATUS_SHORT = {
    "Новое": "🆕 Новое",
    "Принято в работу": "🛠 В работе",
    "Не принято": "🟪 Не принято",
    "К обсуждению": "🟣 К обсужд.",
    "Выполнено": "✅ Выполнено",
}

STATUS_ORDER = ["🆕 Новое", "🛠 В работе", "🟪 Не принято",
                "🟣 К обсужд.", "✅ Выполнено"]


def _strip_prefix(display_value: str) -> str:
    if not display_value:
        return ""
    for prefix in CAT_PREFIX.values():
        if display_value.startswith(prefix):
            return display_value[len(prefix):].strip()
    return display_value.strip()


def _add_prefix(raw_value: str):
    if not raw_value or pd.isna(raw_value):
        return None
    for cat, prefix in CAT_PREFIX.items():
        if raw_value == cat:
            return f"{prefix} {cat}"
    return raw_value


# ---------------------------------------------------------------------------
#  Справочники
# ---------------------------------------------------------------------------
@st.cache_data(ttl=600, show_spinner=False)
def _load_discipline_options() -> dict[str, str]:
    with get_conn() as conn:
        codes = [r["discipline"] for r in conn.execute(
            "SELECT DISTINCT discipline FROM documents "
            "WHERE discipline IS NOT NULL ORDER BY discipline")]
    return {c: f"{c} — {discipline_name(c)}" for c in codes}


@st.cache_data(ttl=600, show_spinner=False)
def _load_section_options(disciplines: tuple = ()) -> dict[str, str]:
    with get_conn() as conn:
        if disciplines:
            ph = ",".join("?" * len(disciplines))
            rows = conn.execute(f"""
                SELECT DISTINCT section FROM documents
                WHERE section IS NOT NULL AND section <> ''
                  AND discipline IN ({ph})
                ORDER BY section
            """, tuple(disciplines)).fetchall()
        else:
            rows = conn.execute("""
                SELECT DISTINCT section FROM documents
                WHERE section IS NOT NULL AND section <> ''
                ORDER BY section
            """).fetchall()
    return {r["section"]: r["section"] for r in rows}


@st.cache_data(ttl=600, show_spinner=False)
def _load_kit_options(disciplines: tuple = (),
                      sections: tuple = ()) -> dict[str, str]:
    with get_conn() as conn:
        where = ["c.code IS NOT NULL"]
        params: list = []

        if disciplines:
            where.append(f"c.discipline IN ({','.join('?' * len(disciplines))})")
            params += list(disciplines)

        if sections:
            ph = ",".join("?" * len(sections))
            where.append(f"""c.code IN (
                SELECT DISTINCT complex FROM documents
                WHERE section IN ({ph})
                  AND complex IS NOT NULL
            )""")
            params += list(sections)

        rows = conn.execute(f"""
            SELECT c.code, c.name
            FROM complexes c
            WHERE {' AND '.join(where)}
            ORDER BY c.code
        """, tuple(params)).fetchall()

    result = {}
    for r in rows:
        code = r["code"]
        name = (r["name"] or "").strip()
        result[code] = f"{code} — {name}" if name else code
    return result


# ---------------------------------------------------------------------------
#  Единый срез наших замечаний
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_active_df() -> pd.DataFrame:
    """Замечания на нашей стороне (holder == "ours")."""
    from vitro.ui.deadlines import _load_all_categorized
    df = _load_all_categorized()
    if df.empty:
        return df
    return df[df["holder"] == "ours"].copy()


# ---------------------------------------------------------------------------
#  Прогресс категоризации
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_categorization_progress() -> dict:
    active = _load_active_df()
    total = len(active)

    if total > 0:
        done = active[
            active["category"].notna()
            & (active["category"].astype(str).str.strip() != "")
        ].shape[0]
    else:
        done = 0

    left = total - done

    with get_conn() as conn:
        closed = conn.execute(
            "SELECT COUNT(*) FROM comments WHERE status = 'Закрыто'"
        ).fetchone()[0]
        annulled = conn.execute(
            "SELECT COUNT(*) FROM comments WHERE status = 'Аннулировано'"
        ).fetchone()[0]

    return {
        "total": total,
        "done": done,
        "left": left,
        "closed": closed,
        "annulled": annulled,
    }


# ---------------------------------------------------------------------------
#  Распределение по 4 категориям
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_overall_distribution() -> pd.DataFrame:
    active = _load_active_df()
    if active.empty:
        return pd.DataFrame(columns=["Категория", "Количество"])

    counts = active["category"].fillna("Без категории").replace(
        "", "Без категории").value_counts()

    rows = []
    for cat in CAT_OPTIONS_RAW:
        rows.append({
            "Категория": cat,
            "Количество": int(counts.get(cat, 0)),
        })
    uncat = int(counts.get("Без категории", 0))
    if uncat > 0:
        rows.append({
            "Категория": "Без категории",
            "Количество": uncat,
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
#  Кросс-таблица
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_category_by_status() -> pd.DataFrame:
    active = _load_active_df()
    if active.empty:
        return pd.DataFrame()

    df = active.copy()
    df["Категория"] = df["category"].fillna("Без категории").replace(
        "", "Без категории")
    df["Статус"] = df["status"].map(STATUS_SHORT).fillna(df["status"])

    pivot = df.pivot_table(
        index="Категория",
        columns="Статус",
        values="id",
        aggfunc="count",
        fill_value=0,
    )

    row_order = CAT_OPTIONS_RAW + ["Без категории"]
    pivot = pivot.reindex([r for r in row_order if r in pivot.index])

    col_order = [c for c in STATUS_ORDER if c in pivot.columns]
    pivot = pivot[col_order]

    pivot["Итого"] = pivot.sum(axis=1)

    totals_row = pivot.sum(axis=0).to_frame().T
    totals_row.index = ["Итого"]
    pivot = pd.concat([pivot, totals_row])

    return pivot.reset_index().rename(columns={"index": "Категория"})


# ---------------------------------------------------------------------------
#  По дисциплинам
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_by_discipline() -> pd.DataFrame:
    active = _load_active_df()
    if active.empty:
        return pd.DataFrame()

    df = active.copy()
    df["Категория"] = df["category"].fillna("Без категории").replace(
        "", "Без категории")

    by_disc = df.groupby(["discipline", "Категория"]).size() \
                .reset_index(name="Количество")
    by_disc = by_disc.rename(columns={"discipline": "Код"})
    return by_disc


# ---------------------------------------------------------------------------
#  Активность специалистов
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_user_activity() -> pd.DataFrame:
    with get_conn() as conn:
        return pd.read_sql("""
            SELECT
                user AS "Специалист",
                COUNT(*) AS "Изменений",
                COUNT(DISTINCT comment_id) AS "Уникальных замечаний",
                MIN(timestamp) AS "Первое",
                MAX(timestamp) AS "Последнее"
            FROM users_activity
            WHERE user IS NOT NULL AND user <> ''
            GROUP BY user
            ORDER BY "Изменений" DESC
        """, conn)


# ---------------------------------------------------------------------------
#  Анализ
# ---------------------------------------------------------------------------
def _render_analysis():
    # =====================================================================
    #  Прогресс категоризации
    # =====================================================================
    st.markdown("### 🎯 Прогресс категоризации")
    st.caption(
        "Считается **только по замечаниям на нашей стороне** "
        "(без листов A и аннулированных)."
    )

    prog = _load_categorization_progress()
    total = prog["total"]
    done = prog["done"]
    left = prog["left"]
    pct = round(done / total * 100, 1) if total else 0

    c1, c2, c3, c4 = st.columns(4)
    c1.metric(
        "Замечаний у АТП ТЛП",
        f"{total:,}".replace(",", " "),
        help="Замечания на нашей стороне (без A и аннулированных).",
    )
    c2.metric("Разобрано", f"{done:,}".replace(",", " "))
    c3.metric("Осталось", f"{left:,}".replace(",", " "))
    c4.metric(
        "Прогресс",
        f"{pct}%",
        help="Доля наших замечаний с категорией.",
    )

    cc1, cc2 = st.columns(2)
    cc1.metric(
        "Закрыто (справочно)",
        f"{prog['closed']:,}".replace(",", " "),
        help="Замечания со статусом «Закрыто». В прогресс не входят.",
    )
    cc2.metric(
        "Аннулировано (справочно)",
        f"{prog['annulled']:,}".replace(",", " "),
        help="Замечания, снятые заказчиком. В прогресс не входят.",
    )

    st.progress(pct / 100)

    with get_conn() as conn:
        recent = conn.execute("""
            SELECT COUNT(*) FROM users_activity
            WHERE timestamp >= datetime('now', '-7 days')
        """).fetchone()[0]

    if recent > 0 and left > 0:
        per_day = recent / 7
        days_left = int(left / per_day) if per_day > 0 else 0
        eta = datetime.now() + timedelta(days=days_left)

        c1, c2, c3 = st.columns(3)
        c1.metric(
            "Темп (7 дн.)",
            f"{per_day:.0f}/день".replace(",", " "),
            help="Среднее число изменений категорий в день за неделю.",
        )
        c2.metric(
            "Осталось дней",
            f"{days_left:,}".replace(",", " "),
            help="При текущем темпе.",
        )
        c3.metric("Прогноз завершения", eta.strftime("%d.%m.%Y"))
    elif left == 0:
        st.success("🎉 Все замечания категоризированы!")
    else:
        st.info(
            "📌 Нет данных за последние 7 дней. Начните категоризировать — "
            "появится прогноз."
        )

    st.divider()

    # =====================================================================
    #  Кросс-таблица: категории × статусы
    # =====================================================================
    st.markdown("### 📊 Категории × 5 статусов")
    st.caption(
        "Какие категории преобладают в каких статусах."
    )

    cross = _load_category_by_status()
    if not cross.empty:
        st.dataframe(
            cross,
            use_container_width=True,
            hide_index=True,
            column_config={
                "Категория": st.column_config.TextColumn(
                    "Категория", width="large"),
            },
        )
    else:
        st.info("Нет замечаний на нашей стороне.")

    st.divider()

    # =====================================================================
    #  Распределение
    # =====================================================================
    st.markdown("### 🏷 Распределение по категориям")

    col1, col2 = st.columns(2)

    with col1:
        dist = _load_overall_distribution()
        if not dist.empty:
            fig = px.pie(
                dist, names="Категория", values="Количество", hole=0.45,
                color="Категория",
                color_discrete_map=CAT_COLORS,
                title="По категориям",
            )
            st.plotly_chart(fig, use_container_width=True)

    with col2:
        by_disc = _load_by_discipline()
        if not by_disc.empty:
            fig = px.bar(
                by_disc, x="Код", y="Количество", color="Категория",
                barmode="stack",
                color_discrete_map=CAT_COLORS,
                title="По дисциплинам",
            )
            fig.update_layout(
                legend_title_text="",
                xaxis_tickangle=-45,
                height=400,
            )
            st.plotly_chart(fig, use_container_width=True)

    st.divider()


    # =====================================================================
    #  Активность
    # =====================================================================
    st.markdown("### 🎯 Активность специалистов по категоризации")
    activity = _load_user_activity()
    if not activity.empty:
        st.dataframe(activity, use_container_width=True, hide_index=True)

        fig = px.bar(
            activity.sort_values("Изменений", ascending=True),
            x="Изменений", y="Специалист", orientation="h",
            text="Изменений",
            color_discrete_sequence=["#64B5F6"],
        )
        fig.update_traces(textposition="outside")
        fig.update_layout(height=max(300, 30 * len(activity)))
        st.plotly_chart(fig, use_container_width=True)
    else:
        st.info("Пока никто не назначал категории.")


# ---------------------------------------------------------------------------
#  Визуальная группировка
# ---------------------------------------------------------------------------
def _collapse_repeats(df: pd.DataFrame, cols: list) -> pd.DataFrame:
    out = df.copy()
    key_prev = None
    for idx, row in out.iterrows():
        key = tuple(row[c] for c in cols)
        if key == key_prev:
            for c in cols:
                out.at[idx, c] = ""
        key_prev = key
    return out


# ---------------------------------------------------------------------------
#  Редактор
# ---------------------------------------------------------------------------
def _render_editor():
    st.markdown("### 📝 Работа с замечаниями")
    st.caption(
        "Редактор показывает **только замечания на нашей стороне** "
        "(Новое, Принято в работу, Не принято, К обсуждению). "
        "Закрытые, аннулированные и замечания к листам A "
        "сюда не попадают."
    )

    st.markdown("""
    <style>
    div[data-testid="stDataFrame"] div[role="gridcell"] {
        white-space: normal !important;
        word-wrap: break-word !important;
        line-height: 1.25 !important;
    }
    div[data-testid="stDataFrame"] div[role="row"] {
        min-height: 56px !important;
    }
    </style>
    """, unsafe_allow_html=True)

    c_show, c_only = st.columns([2, 1])

    with c_show:
        show_waiting = st.checkbox(
            "🔵 Показать «Выполнено» (ждут заказчика)",
            value=False,
            key="editor_show_waiting",
            help="Если выключено — в редакторе только 4 активных статуса. "
                 "Если включено — добавляются замечания со статусом "
                 "«Выполнено» (АТП ТЛП ответил, ждём рассмотрения).",
        )

    with c_only:
        only_uncategorized = st.checkbox(
            "Только без категории",
            value=False,
            key="editor_only_uncat",
        )

    disc_options = _load_discipline_options()

    with st.expander("🎛 Фильтры", expanded=False):
        c1, c2, c3 = st.columns(3)

        with c1:
            sel_disc_labels = st.multiselect(
                "Дисциплина",
                options=list(disc_options.values()),
                placeholder="Все дисциплины",
                key="editor_disc",
            )
            sel_disc = [code for code, label in disc_options.items()
                        if label in sel_disc_labels]

        section_options = _load_section_options(
            tuple(sel_disc) if sel_disc else ())

        with c2:
            if section_options:
                sel_section_labels = st.multiselect(
                    "Раздел",
                    options=list(section_options.values()),
                    placeholder="Все разделы",
                    key="editor_section",
                )
                sel_section = [code for code, label in section_options.items()
                               if label in sel_section_labels]
            else:
                sel_section = []
                st.multiselect(
                    "Раздел", options=[],
                    placeholder="Разделы не применимы",
                    disabled=True,
                    key="editor_section_empty",
                )

        kit_options = _load_kit_options(
            tuple(sel_disc) if sel_disc else (),
            tuple(sel_section) if sel_section else (),
        )

        with c3:
            sel_kit_labels = st.multiselect(
                "Комплект",
                options=list(kit_options.values()),
                placeholder="Все комплекты",
                key="editor_kit",
            )
            sel_kit = [code for code, label in kit_options.items()
                       if label in sel_kit_labels]

        if show_waiting:
            status_options = [
                "Новое", "Принято в работу", "Не принято",
                "К обсуждению", "Выполнено",
            ]
        else:
            status_options = [
                "Новое", "Принято в работу",
                "Не принято", "К обсуждению",
            ]

        c4, c5 = st.columns(2)
        with c4:
            sel_status = st.multiselect(
                "Статус замечания",
                options=status_options,
                default=status_options,
                placeholder="Все доступные статусы",
                key="editor_status",
            )

        with c5:
            limit = st.number_input(
                "Лимит строк",
                min_value=50, max_value=2000,
                value=300, step=50, key="editor_limit",
            )

    df = load_remarks_for_editor(
        disciplines=tuple(sel_disc) if sel_disc else None,
        sections=tuple(sel_section) if sel_section else None,
        kits=tuple(sel_kit) if sel_kit else None,
        statuses=tuple(sel_status) if sel_status else None,
        only_uncategorized=only_uncategorized,
        show_waiting=show_waiting,
        limit=int(limit),
    )

    if df.empty:
        st.info("Нет замечаний по заданным фильтрам.")
        return

    st.caption(
        f"Загружено **{len(df)}** строк. "
        f"Пустые ячейки в колонках «Дисциплина», «Раздел», «Комплект», "
        f"«Лист», «Название листа» означают, что значение совпадает "
        f"с ячейкой выше."
    )

    df["category"] = df["category"].apply(_add_prefix)

    group_cols = ["discipline", "section", "complex", "sheet", "sheet_name"]
    view = _collapse_repeats(df, group_cols)

    desired_order = [
        "id", "discipline", "section", "complex",
        "sheet", "sheet_name", "comment",
        "api_status", "author", "created",
        "category", "category_user", "category_date",
    ]
    view = view[[c for c in desired_order if c in view.columns]]

    edited = st.data_editor(
        view,
        column_config={
            "id":               st.column_config.NumberColumn(
                                    "ID", disabled=True, width="small"),
            "discipline":       st.column_config.TextColumn(
                                    "Дисциплина", disabled=True, width="small"),
            "section":          st.column_config.TextColumn(
                                    "Раздел", disabled=True, width="small"),
            "complex":          st.column_config.TextColumn(
                                    "Комплект", disabled=True, width="medium"),
            "sheet":            st.column_config.TextColumn(
                                    "Лист", disabled=True, width="medium"),
            "sheet_name":       st.column_config.TextColumn(
                                    "Название листа", disabled=True, width="large"),
            "comment":          st.column_config.TextColumn(
                                    "Замечание", disabled=True, width="large"),
            "api_status":       st.column_config.TextColumn(
                                    "Текущий статус замечания", disabled=True),
            "author":           st.column_config.TextColumn(
                                    "Автор", disabled=True),
            "created":          st.column_config.TextColumn(
                                    "Создано", disabled=True, width="small"),
            "category":         st.column_config.SelectboxColumn(
                                    "Категория",
                                    options=CAT_OPTIONS_DISPLAY,
                                    required=False,
                                    help="Выберите категорию"),
            "category_user":    st.column_config.TextColumn(
                                    "Кто", disabled=True, width="small"),
            "category_date":    st.column_config.TextColumn(
                                    "Когда", disabled=True, width="small"),
            "category_version": None,
        },
        disabled=["id", "discipline", "section", "complex", "sheet",
                  "sheet_name", "comment", "api_status", "author",
                  "created", "category_user", "category_date"],
        hide_index=True,
        use_container_width=True,
        key="categorization_editor",
    )

    col_save, _ = st.columns([1, 3])
    with col_save:
        save_clicked = st.button("💾 Сохранить изменения", type="primary",
                                 use_container_width=True)

    if save_clicked:
        user = st.session_state.get("user", "инженер")
        saved, conflicts = 0, []

        for _, row in edited.iterrows():
            orig = df[df["id"] == row["id"]].iloc[0]

            new_cat_display = None if pd.isna(row["category"]) else str(row["category"])
            new_cat = _strip_prefix(new_cat_display) if new_cat_display else None

            old_cat_display = None if pd.isna(orig["category"]) else str(orig["category"])
            old_cat = _strip_prefix(old_cat_display) if old_cat_display else None

            if new_cat == old_cat:
                continue
            if new_cat is None:
                new_cat = ""

            ok, msg = update_category_safe(
                int(row["id"]), new_cat, user=user,
                expected_version=int(orig["category_version"] or 0),
            )
            if ok:
                saved += 1
            else:
                conflicts.append(f"ID {row['id']}: {msg}")

        if saved:
            st.success(f"✅ Сохранено: {saved} строк(и)")
            st.cache_data.clear()
            st.rerun()

        if conflicts:
            st.warning("⚠️ Обнаружены конфликты:")
            for c in conflicts[:20]:
                st.write(f"- {c}")
            if len(conflicts) > 20:
                st.write(f"... и ещё {len(conflicts) - 20}")


# ---------------------------------------------------------------------------
#  Импорт из Excel
# ---------------------------------------------------------------------------
def _normalize_category(value) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    s = str(value).strip()
    if not s:
        return None

    s = _strip_prefix(s)

    if s in CAT_OPTIONS_RAW:
        return s

    s_low = s.lower()
    for cat in CAT_OPTIONS_RAW:
        if cat.lower() == s_low:
            return cat
    if "принято" in s_low and "не" not in s_low.split("принято")[0][-5:]:
        if "не принято" in s_low:
            return CAT_4
        return CAT_1
    if "формальн" in s_low:
        return CAT_2
    if "доп" in s_low and "треб" in s_low:
        return CAT_3
    if "не принято" in s_low:
        return CAT_4

    return None


def _parse_excel_file(uploaded_file) -> pd.DataFrame:
    from openpyxl import load_workbook

    rows = []

    try:
        data = uploaded_file.read()
        uploaded_file.seek(0)
        wb = load_workbook(
            io.BytesIO(data),
            data_only=True,
            read_only=True,
        )
    except Exception as e:
        st.error(f"Не удалось открыть файл {uploaded_file.name}: {e}")
        return pd.DataFrame()

    SKIP_SHEETS = {
        "Ведомость листов",
        "Управление",
    }

    sheets_with_data = []

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        if sheet_name.startswith("Сводка") or sheet_name.startswith("Свод"):
            continue

        ws = wb[sheet_name]

        if ws.sheet_state != "visible":
            continue

        header_row = None
        header_idx = None

        def _norm(s):
            if s is None:
                return ""
            s = str(s)
            s = s.replace("\xa0", " ")
            s = s.replace("\u200b", "")
            s = " ".join(s.split())
            return s.strip()

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 15:
                break
            row_str = [_norm(c) for c in row]
            has_id = any(c.upper() in ("ИД", "ID") for c in row_str)
            has_cat = any("катег" in c.lower() for c in row_str)
            if has_id and has_cat:
                header_row = row_str
                header_idx = i
                break

        if header_row is None:
            continue

        id_col = None
        cat_col = None

        for j, h in enumerate(header_row):
            h_norm = _norm(h).upper()
            h_low = _norm(h).lower()
            if h_norm in ("ИД", "ID"):
                id_col = j
            if "катег" in h_low and cat_col is None:
                cat_col = j

        if id_col is None or cat_col is None:
            continue

        sheet_rows = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i <= header_idx:
                continue
            if row is None or len(row) <= max(id_col, cat_col):
                continue

            id_val = row[id_col]
            cat_val = row[cat_col]

            if id_val is None:
                continue
            if cat_val is None or str(cat_val).strip() == "":
                continue

            try:
                cid = int(id_val)
            except (ValueError, TypeError):
                continue

            cat_str = str(cat_val).strip()
            cat_norm = _normalize_category(cat_str)

            rows.append({
                "id": cid,
                "category_raw": cat_str,
                "category": cat_norm,
                "source_sheet": sheet_name,
                "source_file": uploaded_file.name,
            })
            sheet_rows += 1

        if sheet_rows > 0:
            sheets_with_data.append((sheet_name, sheet_rows))

    wb.close()

    if sheets_with_data:
        total = sum(n for _, n in sheets_with_data)
        st.caption(
            f"Прочитано листов: {len(sheets_with_data)}, "
            f"всего строк: {total}"
        )

    return pd.DataFrame(rows)


def _render_import():
    st.markdown("### 📥 Массовая загрузка категорий из Excel")

    user_default = st.session_state.get("user", "").strip()
    if not user_default or user_default == "инженер":
        st.warning(
            "⚠️ **Введите имя в сайдбаре слева** — оно будет записано "
            "как автор категорий. Без имени импорт заблокирован."
        )

    st.caption(
        "Перетащите один или несколько Excel-файлов. Мы читаем "
        "**только видимые листы** и берём из каждого колонки "
        "**«ИД»** и **«Категория замечания»**."
    )

    uploaded = st.file_uploader(
        "Файлы Excel",
        type=["xlsx", "xls", "xlsm", "xlsb"],
        accept_multiple_files=True,
        key="cat_import_files",
    )

    if not uploaded:
        st.info("Файлы не выбраны.")
        return

    st.markdown(f"**Получено файлов:** {len(uploaded)}")

    with st.expander("🔍 Предпросмотр распознанных категорий", expanded=True):
        all_rows = []
        file_stats = []

        for f in uploaded:
            df_file = _parse_excel_file(f)
            n_rows = len(df_file)
            n_sheets = df_file["source_sheet"].nunique() if not df_file.empty else 0

            file_stats.append({
                "Файл": f.name,
                "Строк": n_rows,
                "Листов с данными": n_sheets,
            })

            if not df_file.empty:
                all_rows.append(df_file)

        if file_stats:
            st.markdown("**Диагностика файлов:**")
            st.dataframe(pd.DataFrame(file_stats),
                         use_container_width=True, hide_index=True)

        if not all_rows:
            st.warning("Ни в одном файле не найдено колонок «ИД» и «Категория».")
            return

        combined = pd.concat(all_rows, ignore_index=True)
        combined_unique = combined.drop_duplicates(subset=["id"], keep="last")

        total_rows = len(combined)
        unique_rows = len(combined_unique)
        unknown_cat = combined_unique["category"].isna().sum()
        n_sheets = combined_unique["source_sheet"].nunique()

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Всего строк", total_rows)
        c2.metric("Уникальных ID", unique_rows)
        c3.metric("Нераспознанных", int(unknown_cat))
        c4.metric("Листов", n_sheets)

        dist = combined_unique["category"].value_counts(dropna=False).reset_index()
        dist.columns = ["Категория", "Количество"]
        st.dataframe(dist, use_container_width=True, hide_index=True)

        if unknown_cat > 0:
            st.warning("Значения, которые не удалось распознать:")
            unknown_df = combined_unique[combined_unique["category"].isna()][
                ["id", "category_raw", "source_file", "source_sheet"]
            ].head(20)
            st.dataframe(unknown_df, use_container_width=True, hide_index=True)

    st.divider()
    st.markdown("#### 👤 Автор импорта")

    import_user = st.text_input(
        "Кто выполняет импорт (имя для записи)",
        value=user_default if user_default != "инженер" else "",
        placeholder="Фамилия Имя Отчество",
        key="cat_import_user",
    )

    force_overwrite = st.checkbox(
        "🔄 Перезаписать существующие категории",
        value=False,
        key="cat_import_force",
        help="Если включено — импорт перезапишет категории.",
    )

    apply = st.button(
        "✅ Применить категории к базе",
        type="primary",
        use_container_width=True,
        disabled=not import_user.strip(),
    )

    if not import_user.strip():
        st.info("Введите имя автора импорта, чтобы разблокировать кнопку.")

    if apply:
        if not all_rows:
            st.error("Нет данных для применения.")
            return

        ids = combined_unique["id"].astype(int).tolist()
        with get_conn() as conn:
            placeholders = ",".join("?" * len(ids))
            existing_rows = conn.execute(f"""
                SELECT id, COALESCE(category_version, 0) AS v
                FROM comments WHERE id IN ({placeholders})
            """, tuple(ids)).fetchall()
        existing = {r["id"]: r["v"] for r in existing_rows}

        saved = 0
        skipped_missing = 0
        skipped_unknown = 0
        conflicts = []

        for _, row in combined_unique.iterrows():
            cid = int(row["id"])
            cat = row["category"]

            if cat is None or pd.isna(cat):
                skipped_unknown += 1
                continue
            if cid not in existing:
                skipped_missing += 1
                continue

            row_user = import_user.strip()

            if force_overwrite:
                ok, msg = update_category_force(cid, cat, user=row_user)
            else:
                ok, msg = update_category_safe(
                    cid, cat, user=row_user,
                    expected_version=existing[cid],
                )

            if ok:
                saved += 1
                existing[cid] = existing[cid] + 1
            else:
                conflicts.append(f"ID {cid}: {msg}")

        st.success(
            f"✅ Обновлено: **{saved}** · "
            f"Пропущено (нет в БД): **{skipped_missing}** · "
            f"Нераспознанных: **{skipped_unknown}** · "
            f"Конфликтов: **{len(conflicts)}**"
        )

        if conflicts:
            with st.expander(f"⚠️ Конфликты ({len(conflicts)})", expanded=True):
                for c in conflicts[:50]:
                    st.write(f"- {c}")

        if saved:
            st.cache_data.clear()
            st.success(f"✅ Обновлено {saved} записей.")
            st.rerun()


# ---------------------------------------------------------------------------
#  Точка входа
# ---------------------------------------------------------------------------
def render():
    st.header("🏷 Категории")
    tab_analysis, tab_editor, tab_import = st.tabs([
        "📈 Анализ",
        "📝 Работа с замечаниями",
        "📥 Импорт из Excel",
    ])

    with tab_analysis:
        _render_analysis()
    with tab_editor:
        _render_editor()
    with tab_import:
        _render_import()
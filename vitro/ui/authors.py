# vitro/ui/authors.py
"""
👤 Авторы — аналитика по инженерам заказчика.

Данные — из единого источника `_load_all_categorized` (deadlines.py).
Логика подсчёта — как на вкладке «Сводка по проекту».

3 под-вкладки:
  📊 Аналитика авторов — KPI, таблица авторов, детали автора.
  🔵 Ждут заказчика    — редактор категорий.
  📄 К листам A и аннулированным — список замечаний (просмотр).
"""

import io
from datetime import datetime

import openpyxl
import pandas as pd
import plotly.express as px
import streamlit as st

from vitro.sqlite_db import get_conn, update_category_safe
from vitro.disciplines import discipline_name
from vitro.ui._auth import ensure_user
from vitro.ui._filters import (
    REMARK_STATUS_OPTIONS,
    SHEET_GROUP_OPTIONS,
    _doc_status_group,
    _filter_by_marka,
)
from vitro.ui._urls import build_vitro_comment_url


# ---- Категории замечаний ----
CAT_1 = "Принято/корректное"
CAT_2 = "Формальное/нет влияния на СМР"
CAT_3 = "Доп.требование/отсутствует в ТЗ"
CAT_4 = "Не принято/нарушение ТНПА"

CAT_PREFIX = {
    CAT_1: "🟢",
    CAT_2: "🟡",
    CAT_3: "🔵",
    CAT_4: "🔴",
}

CAT_OPTIONS_DISPLAY = [
    f"{CAT_PREFIX[c]} {c}" for c in [CAT_1, CAT_2, CAT_3, CAT_4]
]

CAT_COLORS = {
    "Принято":       "#C6EFCE",
    "Формальное":    "#FFEB9C",
    "Доп.треб.":     "#BDD7EE",
    "Не принято":    "#FFC7CE",
    "Без категории": "#D9D9D9",
}

_CAT_PREFIX_MAP = {
    CAT_1: f"🟢 {CAT_1}",
    CAT_2: f"🟡 {CAT_2}",
    CAT_3: f"🔵 {CAT_3}",
    CAT_4: f"🔴 {CAT_4}",
}

# ---- Расположение (4-й блок шифра) ----
BLOCK_OPTIONS = {
    "Корпус 1":                    "1",
    "Корпус 2":                    "2",
    "Общие":                       "0",
    "Стилобат":                    "С",
    "Газовая котельная":           "ГК",
    "Генплан":                     "ГП",
    "Автомобильные дороги (УДС)":  "А",
}

ALL_AUTHORS = "— Все авторы —"


def _strip_prefix(display_value: str) -> str:
    if not display_value:
        return ""
    for prefix in CAT_PREFIX.values():
        if display_value.startswith(prefix):
            return display_value[len(prefix):].strip()
    return display_value.strip()


def _extract_4th_block(complex_code) -> str:
    if not complex_code:
        return ""
    parts = str(complex_code).split("-")
    if len(parts) >= 4:
        return parts[3].strip().upper()
    return ""


def _filter_by_blocks(df: pd.DataFrame,
                       block_codes: tuple = ()) -> pd.DataFrame:
    if df.empty or not block_codes:
        return df
    blocks = df["complex"].apply(_extract_4th_block)
    return df[blocks.isin(block_codes)]


def _collapse_repeats(df: pd.DataFrame,
                       cols: list) -> pd.DataFrame:
    """Скрывает повторяющиеся значения в колонках."""
    out = df.copy()
    key_prev = None
    for idx, row in out.iterrows():
        key = tuple(row[c] for c in cols)
        if key == key_prev:
            for c in cols:
                out.at[idx, c] = ""
        key_prev = key
    return out


# ---------------------------------------------------------------------------
#  Единый источник
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_all_categorized() -> pd.DataFrame:
    from vitro.ui.deadlines import _load_all_categorized as _src
    return _src()


# ---------------------------------------------------------------------------
#  Справочники
# ---------------------------------------------------------------------------
@st.cache_data(ttl=600, show_spinner=False)
def _load_discipline_options() -> dict[str, str]:
    with get_conn() as conn:
        codes = [r["discipline"] for r in conn.execute(
            "SELECT DISTINCT discipline FROM documents "
            "WHERE discipline IS NOT NULL ORDER BY discipline")]
    return {c: f"{c} — {discipline_name(c)}" for c in codes}


@st.cache_data(ttl=600, show_spinner=False)
def _load_section_options(disciplines: tuple = ()) -> dict[str, str]:
    with get_conn() as conn:
        if disciplines:
            ph = ",".join("?" * len(disciplines))
            rows = conn.execute(f"""
                SELECT DISTINCT section FROM documents
                WHERE section IS NOT NULL AND section <> ''
                  AND discipline IN ({ph}) ORDER BY section
            """, tuple(disciplines)).fetchall()
        else:
            rows = conn.execute("""
                SELECT DISTINCT section FROM documents
                WHERE section IS NOT NULL AND section <> ''
                ORDER BY section
            """).fetchall()
    return {r["section"]: r["section"] for r in rows}


@st.cache_data(ttl=600, show_spinner=False)
def _load_kit_options(disciplines: tuple = (),
                       sections: tuple = (),
                       block_codes: tuple = ()) -> dict[str, str]:
    with get_conn() as conn:
        where = ["c.code IS NOT NULL"]
        params: list = []
        if disciplines:
            where.append(f"c.discipline IN ({','.join('?' * len(disciplines))})")
            params += list(disciplines)
        if sections:
            ph = ",".join("?" * len(sections))
            where.append(f"""c.code IN (
                SELECT DISTINCT complex FROM documents
                WHERE section IN ({ph}) AND complex IS NOT NULL)""")
            params += list(sections)
        rows = conn.execute(f"""
            SELECT c.code, c.name FROM complexes c
            WHERE {' AND '.join(where)} ORDER BY c.code
        """, tuple(params)).fetchall()

    result = {r["code"]: f"{r['code']} — {r['name']}" if r["name"] else r["code"]
              for r in rows}

    # Фильтр по расположению (4-й блок) — в Python
    if block_codes:
        result = {code: label for code, label in result.items()
                  if _extract_4th_block(code) in block_codes}
    return result


# ---------------------------------------------------------------------------
#  Фильтрация по иерархии
# ---------------------------------------------------------------------------
def _filter_by_hierarchy(df: pd.DataFrame,
                          disciplines: tuple = (),
                          sections: tuple = (),
                          kits: tuple = (),
                          sheet_groups: tuple = (),
                          remark_statuses: tuple = (),
                          marka: str = "") -> pd.DataFrame:
    if df.empty:
        return df
    out = df
    if disciplines:
        out = out[out["discipline"].isin(disciplines)]
    if sections:
        out = out[out["section"].isin(sections)]
    if kits:
        out = out[out["complex"].isin(kits)]
    if sheet_groups:
        out = out[
            out["doc_status"].apply(_doc_status_group).isin(sheet_groups)
        ]
    if remark_statuses:
        out = out[out["status"].isin(remark_statuses)]
    if marka:
        out = _filter_by_marka(out, marka)
    return out


# ---------------------------------------------------------------------------
#  KPI
# ---------------------------------------------------------------------------
def _render_kpi(df_all: pd.DataFrame, df_ours: pd.DataFrame):
    if df_all.empty:
        st.info("Нет данных по авторам с учётом фильтров.")
        return

    total = len(df_all)
    n_authors = df_all["author"].nunique()

    cat_counts = (
        df_ours["category"].fillna("Без категории")
        .replace("", "Без категории")
        .value_counts()
    )

    st.markdown("##### 📦 Общий объём")
    st.caption(
        "Источник — Сводка по проекту. «Активных» — все замечания "
        "в 5 статусах. Категории — только по нашим 4 статусам, "
        "без замечаний к листам A и к аннулированным."
    )
    c1, c2 = st.columns(2)
    c1.metric(
        "Авторов", n_authors,
        help="Уникальных авторов замечаний в срезе.",
    )
    c2.metric(
        "Активных замечаний", f"{total:,}".replace(",", " "),
        help="Выгрузка из Витро: 5 статусов — Новое, Принято в работу, "
             "Не принято, К обсуждению, Выполнено. "
             "Включая замечания к листам A и к аннулированным.",
    )

    st.markdown("##### 🏷 Категории замечаний")
    st.caption(
        "Оценка АТП ТЛП ставится только по нашим 4 статусам "
        "(Новое, Принято в работу, Не принято, К обсуждению). "
        "Без замечаний к листам A и к аннулированным."
    )
    c1, c2, c3, c4, c5 = st.columns(5)

    c1.metric("🟢 Принято/корректное",
              f"{int(cat_counts.get(CAT_1, 0)):,}".replace(",", " "))
    c2.metric("🟡 Формальное",
              f"{int(cat_counts.get(CAT_2, 0)):,}".replace(",", " "))
    c3.metric("🔵 Доп.требование",
              f"{int(cat_counts.get(CAT_3, 0)):,}".replace(",", " "))
    c4.metric("🔴 Не принято",
              f"{int(cat_counts.get(CAT_4, 0)):,}".replace(",", " "))
    c5.metric("⚪ Без категории",
              f"{int(cat_counts.get('Без категории', 0)):,}".replace(",", " "))


# ---------------------------------------------------------------------------
#  Агрегация по авторам
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_authors(disciplines: tuple = (), sections: tuple = (),
                   kits: tuple = (),
                   sheet_groups: tuple = (),
                   remark_statuses: tuple = (),
                   marka: str = "") -> pd.DataFrame:
    df = _load_all_categorized()
    df = _filter_by_hierarchy(df, disciplines, sections, kits,
                               sheet_groups, remark_statuses, marka)

    if df.empty:
        return pd.DataFrame()

    # Просрочено им (заказчик не рассмотрел > 10 р.д.)
    customers = df[df["holder"] == "customer"]
    customers_overdue = customers[
        customers["customer_bucket"].isin(["10–30", "30–90", ">90"])
    ]
    overdue_by_author = customers_overdue.groupby("author").size()

    # Наши замечания
    ours = df[df["holder"] == "ours"].copy()
    if ours.empty:
        return pd.DataFrame()

    ours["Категория"] = (
        ours["category"].fillna("Без категории")
        .replace("", "Без категории")
    )

    rows = []
    for author, sub in ours.groupby("author"):
        if not author:
            continue

        total = len(sub)
        cat_counts = sub["Категория"].value_counts()

        cnt_1 = int(cat_counts.get(CAT_1, 0))
        cnt_2 = int(cat_counts.get(CAT_2, 0))
        cnt_3 = int(cat_counts.get(CAT_3, 0))
        cnt_4 = int(cat_counts.get(CAT_4, 0))
        cnt_none = int(cat_counts.get("Без категории", 0))
        categorized = cnt_1 + cnt_2 + cnt_3 + cnt_4

        pct_accepted = (
            round(cnt_1 / categorized * 100, 1)
            if categorized > 0 else 0.0
        )

        n_customer_overdue = int(overdue_by_author.get(author, 0))

        rows.append({
            "author": author,
            "total": total,
            "customer_overdue": n_customer_overdue,
            "cat_1": cnt_1,
            "cat_2": cnt_2,
            "cat_3": cnt_3,
            "cat_4": cnt_4,
            "cat_none": cnt_none,
            "categorized": categorized,
            "pct_accepted": pct_accepted,
        })

    result = pd.DataFrame(rows)
    return result if not result.empty else pd.DataFrame()


# ---------------------------------------------------------------------------
#  Таблица авторов
# ---------------------------------------------------------------------------
def _render_authors_table(df: pd.DataFrame):
    st.markdown("### 📋 Авторы замечаний")
    st.caption(
        "Только замечания, ожидающие ответа АТП ТЛП (без листов A "
        "и аннулированных). "
        "«Просрочено им» — заказчик не рассмотрел наш ответ > 10 р.д. "
        "Категории — оценка АТП ТЛП."
    )

    if df.empty:
        st.info("Нет данных.")
        return

    table = df[[
        "author", "total", "customer_overdue",
        "cat_1", "cat_2", "cat_3", "cat_4",
        "pct_accepted",
    ]].rename(columns={
        "author": "Автор",
        "total": "Всего выдал",
        "customer_overdue": "Просрочено им",
        "cat_1": "🟢 Принято",
        "cat_2": "🟡 Формальное",
        "cat_3": "🔵 Доп.треб.",
        "cat_4": "🔴 Не принято",
        "pct_accepted": "% принятых",
    }).sort_values("Всего выдал", ascending=False)

    st.dataframe(
        table,
        use_container_width=True,
        hide_index=True,
        height=600,
        column_config={
            "% принятых": st.column_config.ProgressColumn(
                "% принятых", min_value=0, max_value=100,
                format="%.1f%%"),
        },
    )


# ---------------------------------------------------------------------------
#  Детали автора
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_author_by_discipline(author: str,
                                sheet_groups: tuple = (),
                                remark_statuses: tuple = (),
                                marka: str = "") -> pd.DataFrame:
    df = _load_all_categorized()
    ours = df[df["holder"] == "ours"]
    if ours.empty:
        return pd.DataFrame()

    if sheet_groups:
        ours = ours[
            ours["doc_status"].apply(_doc_status_group).isin(sheet_groups)
        ]
    if remark_statuses:
        ours = ours[ours["status"].isin(remark_statuses)]
    if marka:
        ours = _filter_by_marka(ours, marka)

    sub = ours[ours["author"] == author].copy()
    if sub.empty:
        return pd.DataFrame()

    sub["Категория"] = (
        sub["category"].fillna("Без категории").replace("", "Без категории")
    )

    rows = []
    for disc, g in sub.groupby("discipline"):
        if not disc:
            continue
        cc = g["Категория"].value_counts()
        rows.append({
            "discipline": disc,
            "n": len(g),
            "cat_1": int(cc.get(CAT_1, 0)),
            "cat_2": int(cc.get(CAT_2, 0)),
            "cat_3": int(cc.get(CAT_3, 0)),
            "cat_4": int(cc.get(CAT_4, 0)),
            "cat_none": int(cc.get("Без категории", 0)),
        })
    return pd.DataFrame(rows).sort_values("n", ascending=False)


@st.cache_data(ttl=3600, show_spinner=False)
def _load_author_comments(author: str,
                          sheet_groups: tuple = (),
                          remark_statuses: tuple = (),
                          marka: str = "") -> pd.DataFrame:
    df = _load_all_categorized()
    ours = df[df["holder"] == "ours"]
    if ours.empty:
        return pd.DataFrame()

    if sheet_groups:
        ours = ours[
            ours["doc_status"].apply(_doc_status_group).isin(sheet_groups)
        ]
    if remark_statuses:
        ours = ours[ours["status"].isin(remark_statuses)]
    if marka:
        ours = _filter_by_marka(ours, marka)

    sub = ours[ours["author"] == author].copy()
    if sub.empty:
        return pd.DataFrame()

    view = sub[[
        "id", "discipline", "complex", "sheet", "comment",
        "status", "doc_status", "created", "category",
        "vitro_guid", "vitro_version",
    ]].copy()

    # Собираем прямую ссылку на замечание в Витро
    view["_vitro_url"] = view.apply(
        lambda r: build_vitro_comment_url(
            r["vitro_guid"], r["vitro_version"], r["id"]
        ),
        axis=1,
    )

    view = view.drop(columns=["vitro_guid", "vitro_version"])

    view = view.rename(columns={
        "id": "ID",
        "discipline": "Дисциплина",
        "complex": "Комплект",
        "sheet": "Лист",
        "comment": "Замечание",
        "status": "Статус замечания",
        "doc_status": "Статус листа",
        "created": "Создано",
        "category": "Категория",
        "_vitro_url": "Витро",
    })
    return view


_CAT_LABELS = {
    CAT_1: "Принято",
    CAT_2: "Формальное",
    CAT_3: "Доп.треб.",
    CAT_4: "Не принято",
    "Без категории": "Без категории",
}

_COMMENTS_RENAME = {
    "id": "ID",
    "discipline": "Дисциплина",
    "complex": "Комплект",
    "sheet": "Лист",
    "comment": "Замечание",
    "status": "Статус замечания",
    "doc_status": "Статус листа",
    "created": "Создано",
    "category": "Категория",
    "_vitro_url": "Витро",
}


def _build_cat_data(df: pd.DataFrame) -> pd.DataFrame:
    """Данные для бублика категорий по переданному срезу."""
    cats = (df["category"].fillna("Без категории")
            .replace("", "Без категории"))
    counts = cats.value_counts()
    return pd.DataFrame([
        {"Категория": label, "Количество": int(counts.get(cat, 0))}
        for cat, label in _CAT_LABELS.items()
    ])


def _build_comments_view(df: pd.DataFrame) -> pd.DataFrame:
    """Таблица замечаний для drilldown — по текущим фильтрам."""
    cols = ["id", "discipline", "complex", "sheet", "comment",
            "status", "doc_status", "created", "category",
            "vitro_guid", "vitro_version"]
    cols = [c for c in cols if c in df.columns]
    view = df[cols].copy()

    view["_vitro_url"] = view.apply(
        lambda r: build_vitro_comment_url(
            r["vitro_guid"], r["vitro_version"], r["id"]
        ),
        axis=1,
    )

    view = view.drop(columns=["vitro_guid", "vitro_version"], errors="ignore")
    return view.rename(columns=_COMMENTS_RENAME)


def _render_drilldown(df: pd.DataFrame, df_ours: pd.DataFrame):
    st.markdown("### 🔍 Детали автора")
    if df.empty:
        st.info("Нет данных.")
        return

    authors = [ALL_AUTHORS] + sorted(
        a for a in df["author"].dropna().unique() if a
    )

    # Сброс «залипшего» выбора, если автор пропал после смены фильтров
    if st.session_state.get("author_drill") not in authors:
        st.session_state["author_drill"] = ALL_AUTHORS

    sel = st.selectbox(
        "Выберите автора", options=authors,
        key="author_drill",
        placeholder="Начните вводить имя...",
    )

    if sel == ALL_AUTHORS:
        sub = df
        chart_df = df_ours
        pie_title = "Категории замечаний — все авторы"
    else:
        sub = df[df["author"] == sel]
        chart_df = df_ours[df_ours["author"] == sel]
        pie_title = "Категории замечаний автора"

    if sub.empty:
        st.info("Нет данных по выбранному автору.")
        return

    n_total = int(sub["total"].sum())
    n_overdue = int(sub["customer_overdue"].sum())
    cnt_1 = int(sub["cat_1"].sum())
    cnt_2 = int(sub["cat_2"].sum())
    cnt_3 = int(sub["cat_3"].sum())
    cnt_4 = int(sub["cat_4"].sum())
    cnt_none = int(sub["cat_none"].sum())
    categorized = cnt_1 + cnt_2 + cnt_3 + cnt_4
    pct_accepted = (round(cnt_1 / categorized * 100, 1)
                    if categorized > 0 else 0.0)

    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Всего выдал", n_total,
              help="Замечания, ожидающие ответа АТП ТЛП.")
    c2.metric("Просрочено им", n_overdue,
              help="Заказчик не рассмотрел наш ответ > 10 р.д.")
    c3.metric("Принято", cnt_1,
              help="Оценка АТП ТЛП: замечание принято как корректное.")
    c4.metric("Не принято", cnt_4,
              help="Оценка АТП ТЛП: замечание не принимаем.")
    c5.metric("% принятых", f"{pct_accepted}%",
              help="Принято / (Принято + Формальное + Доп.треб. + "
                   "Не принято). «Без категории» не учитывается.")

    st.divider()

    col1, col2 = st.columns(2)

    with col1:
        cat_data = _build_cat_data(chart_df)
        cat_data = cat_data[cat_data["Количество"] > 0]
        if cat_data.empty:
            st.caption("Нет данных по категориям.")
        else:
            fig = px.pie(
                cat_data, names="Категория", values="Количество",
                hole=0.45, color="Категория",
                color_discrete_map=CAT_COLORS,
                title=pie_title,
            )
            fig.update_traces(
                textposition="inside", textinfo="percent",
                textfont_size=11,
            )
            fig.update_layout(
                showlegend=True,
                margin=dict(l=10, r=10, t=40, b=10),
                height=320,
            )
            st.plotly_chart(fig, use_container_width=True,
                            key="authors_drill_cat_pie")

    with col2:
        by_disc = (
            chart_df.groupby("discipline").size()
            .reset_index(name="n")
        )
        if by_disc.empty:
            st.caption("Нет данных по дисциплинам.")
        else:
            by_disc["label"] = by_disc["discipline"].apply(
                lambda c: f"{c} — {discipline_name(c)}")
            fig = px.bar(
                by_disc.sort_values("n", ascending=True),
                x="n", y="label", orientation="h", text="n",
                labels={"n": "Замечаний", "label": ""},
                color_discrete_sequence=["#9575CD"],
            )
            fig.update_traces(textposition="outside")
            fig.update_layout(height=max(300, 25 * len(by_disc)))
            st.plotly_chart(fig, use_container_width=True,
                            key="authors_drill_disc_bar")

    st.divider()

    if sel == ALL_AUTHORS:
        st.markdown("**Все замечания — все авторы**")
    else:
        st.markdown(f"**Все замечания автора `{sel}`**")

    if chart_df.empty:
        st.info("Нет замечаний по текущим фильтрам.")
        return

    # Ограничиваем вывод, чтобы браузер не зависал
    MAX_ROWS = 500
    n_details = len(chart_df)
    details = _build_comments_view(chart_df.head(MAX_ROWS))

    if n_details > MAX_ROWS:
        n_details_fmt = f"{n_details:,}".replace(",", " ")
        st.caption(
            f"Показано **{MAX_ROWS}** из **{n_details_fmt}** замечаний."
        )

    if details.empty:
        st.info("У автора нет замечаний.")
        return
    st.dataframe(
        details, use_container_width=True, hide_index=True,
        height=500,
        column_config={
            "Замечание": st.column_config.TextColumn(width="large"),
            "Витро": st.column_config.LinkColumn(
                "Витро",
                display_text="🔗 Открыть",
                width="small",
                help="Открыть замечание в Витро (подсветится на листе)",
            ),
        },
    )


# ---------------------------------------------------------------------------
#  Ждут заказчика — редактор категорий
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_waiting_review(disciplines=None, kits=None,
                          sections=None, authors=None) -> pd.DataFrame:
    df = _load_all_categorized()
    if df.empty:
        return df

    df = df[df["holder"] == "customer"].copy()
    if df.empty:
        return df

    if disciplines:
        df = df[df["discipline"].isin(disciplines)]
    if sections:
        df = df[df["section"].isin(sections)]
    if kits:
        df = df[df["complex"].isin(kits)]
    if authors:
        df = df[df["author"].isin(authors)]

    df = df.copy()
    df["fix_date"] = df["fix_date"].fillna("")
    df["days_waiting"] = df["days_waiting_customer"]

    return df


def _render_waiting_review():
    st.markdown("### 🔵 Ответ дан, ожидают рассмотрения заказчика")
    st.caption(
        "Замечания в статусе **«Выполнено»**. АТП ТЛП ответил — ждём "
        "решения заказчика. Категорию можно поставить или изменить. "
        "Без замечаний к листам A и к аннулированным."
    )

    _df_for_authors = _load_waiting_review()
    if not _df_for_authors.empty:
        _authors = sorted(
            _df_for_authors["author"].dropna().unique().tolist()
        )
    else:
        _authors = []

    disc_options = _load_discipline_options()
    with st.expander("🎛 Фильтры", expanded=False):
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            sel_disc_labels = st.multiselect(
                "Дисциплина", options=list(disc_options.values()),
                placeholder="Все дисциплины", key="wait_disc",
            )
            sel_disc = [c for c, l in disc_options.items()
                         if l in sel_disc_labels]
        section_options = _load_section_options(
            tuple(sel_disc) if sel_disc else ())
        with c2:
            if section_options:
                sel_section_labels = st.multiselect(
                    "Раздел", options=list(section_options.values()),
                    placeholder="Все разделы", key="wait_section",
                )
                sel_section = [c for c, l in section_options.items()
                                if l in sel_section_labels]
            else:
                sel_section = []
                st.multiselect("Раздел", options=[],
                                placeholder="Не применимы",
                                disabled=True, key="wait_section_empty")
        kit_options = _load_kit_options(
            tuple(sel_disc) if sel_disc else (),
            tuple(sel_section) if sel_section else ())
        with c3:
            sel_kit_labels = st.multiselect(
                "Комплект", options=list(kit_options.values()),
                placeholder="Все комплекты", key="wait_kit",
            )
            sel_kit = [c for c, l in kit_options.items()
                        if l in sel_kit_labels]
        with c4:
            sel_author_labels = st.multiselect(
                "Автор (заказчик)", options=_authors,
                placeholder="Все авторы", key="wait_author",
            )
            sel_author = sel_author_labels

    df = _load_waiting_review(
        disciplines=tuple(sel_disc) if sel_disc else None,
        kits=tuple(sel_kit) if sel_kit else None,
        sections=tuple(sel_section) if sel_section else None,
        authors=tuple(sel_author) if sel_author else None,
    )

    if df.empty:
        st.success("🎉 Нет замечаний, ожидающих рассмотрения заказчика.")
        return

    buckets_order = ["≤10", "10–30", "30–90", ">90"]
    counts = {b: int((df["customer_bucket"] == b).sum())
              for b in buckets_order}

    n_in_time = counts["≤10"]
    n_overdue = counts["10–30"] + counts["30–90"] + counts[">90"]
    n_total = n_in_time + n_overdue

    st.markdown("##### 📊 Классификация")

    st.markdown("**🔵 Заказчик должен рассмотреть**")
    st.caption(
        "Мы ответили (статус «Выполнено»), но заказчик ещё не "
        "рассмотрел. Это **его** зона ответственности."
    )
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric(
        "🟢 срок не превышен", f"{n_in_time:,}".replace(",", " "),
        help="Заказчик ещё в пределах 10 рабочих дней.",
    )
    c2.metric(
        "🟡 10–30 р.д.", f"{counts['10–30']:,}".replace(",", " "),
        help="Недавно вышли за SLA.",
    )
    c3.metric(
        "🟠 30–90 р.д.", f"{counts['30–90']:,}".replace(",", " "),
        help="Значительное ожидание — нужен письменный запрос.",
    )
    c4.metric(
        "🔴 >90 р.д.", f"{counts['>90']:,}".replace(",", " "),
        help="Критическое ожидание — эскалация руководству.",
    )
    c5.metric(
        "📊 Итого ждут рассмотрения",
        f"{n_total:,}".replace(",", " "),
        help="Сумма всех замечаний, ожидающих рассмотрения заказчиком.",
    )

    st.divider()

    st.markdown(f"##### 📋 Список ({len(df):,})".replace(",", " "))

    df["category"] = df["category"].map(_CAT_PREFIX_MAP).fillna(df["category"])

    bucket_priority = {">90": 0, "30–90": 1, "10–30": 2, "≤10": 3}
    df["_sort"] = df["customer_bucket"].map(bucket_priority)
    df = df.sort_values(["_sort", "days_waiting"],
                         ascending=[True, False]).drop(columns=["_sort"])

    display_cols = [
        "id", "customer_bucket", "doc_status",
        "discipline", "complex", "sheet",
        "comment", "author", "fix_date", "days_waiting",
        "category", "category_user", "category_date",
        "category_version",
    ]
    display_cols = [c for c in display_cols if c in df.columns]
    view = df[display_cols].copy()

    MAX_ROWS = 500
    view_limited = view.head(MAX_ROWS)

    if len(view) > MAX_ROWS:
        st.caption(
            f"Показаны первые **{MAX_ROWS}** из **{len(view):,}** строк."
            .replace(",", " ")
        )

    edited = st.data_editor(
        view_limited,
        column_config={
            "id": st.column_config.NumberColumn(
                "ID", disabled=True, width="small"),
            "customer_bucket": st.column_config.TextColumn(
                "Бакет ожидания", disabled=True),
            "discipline": st.column_config.TextColumn(
                "Дисц.", disabled=True, width="small"),
            "complex": st.column_config.TextColumn(
                "Комплект", disabled=True, width="medium"),
            "sheet": st.column_config.TextColumn(
                "Лист", disabled=True, width="medium"),
            "doc_status": st.column_config.TextColumn(
                "Статус листа", disabled=True, width="small"),
            "comment": st.column_config.TextColumn(
                "Замечание", disabled=True, width="large"),
            "author": st.column_config.TextColumn("Автор", disabled=True),
            "fix_date": st.column_config.TextColumn(
                "Наш ответ", disabled=True, width="small"),
            "days_waiting": st.column_config.NumberColumn(
                "Ждём (р.д.)", disabled=True, width="small"),
            "category": st.column_config.SelectboxColumn(
                "Категория", options=CAT_OPTIONS_DISPLAY,
                required=False),
            "category_user": st.column_config.TextColumn(
                "Кто", disabled=True, width="small"),
            "category_date": st.column_config.TextColumn(
                "Когда", disabled=True, width="small"),
            "category_version": None,
        },
        disabled=["id", "customer_bucket", "discipline", "complex", "sheet",
                  "comment", "author", "fix_date", "days_waiting",
                  "category_user", "category_date"],
        hide_index=True,
        use_container_width=True,
        height=600,
        key="waiting_review_editor",
    )

    col_save, _ = st.columns([1, 3])

    with col_save:
        if st.button("💾 Сохранить изменения", type="primary",
                      use_container_width=True, key="wait_save"):
            user = ensure_user()
            if user is None:
                st.warning("⚠️ Введите имя в диалоге, чтобы сохранить изменения.")
                return

            saved, conflicts = 0, []

            df_by_id = df.set_index("id", drop=False)

            for _, row in edited.iterrows():
                row_id = int(row["id"])
                if row_id not in df_by_id.index:
                    continue
                orig = df_by_id.loc[row_id]
                new_cat_display = (None if pd.isna(row["category"])
                                    else str(row["category"]))
                new_cat = (_strip_prefix(new_cat_display)
                            if new_cat_display else None)
                old_cat_display = (None if pd.isna(orig["category"])
                                    else str(orig["category"]))
                old_cat = (_strip_prefix(old_cat_display)
                            if old_cat_display else None)

                if new_cat == old_cat:
                    continue
                if new_cat is None:
                    new_cat = ""

                ok, msg = update_category_safe(
                    int(row["id"]), new_cat, user=user,
                    expected_version=int(orig["category_version"] or 0),
                )
                if ok:
                    saved += 1
                else:
                    conflicts.append(f"ID {row['id']}: {msg}")

            if saved:
                st.success(f"✅ Сохранено: {saved} строк(и)")
                _load_waiting_review.clear()
                st.rerun()
            if conflicts:
                st.warning("⚠️ Конфликты:")
                for c in conflicts[:20]:
                    st.write(f"- {c}")


# ---------------------------------------------------------------------------
#  К листам A и аннулированным
# ---------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def _load_customer_a_annulled(disciplines=None, kits=None,
                                sections=None,
                                block_codes=None) -> pd.DataFrame:
    """
    Замечания к листам A и к аннулированным листам.
    holder IN ('customer_a', 'annulled').
    """
    df = _load_all_categorized()
    if df.empty:
        return df

    df = df[df["holder"].isin(["customer_a", "annulled"])].copy()
    if df.empty:
        return df

    if disciplines:
        df = df[df["discipline"].isin(disciplines)]
    if sections:
        df = df[df["section"].isin(sections)]
    if kits:
        df = df[df["complex"].isin(kits)]
    if block_codes:
        df = _filter_by_blocks(df, tuple(block_codes))

    return df


_CA_COLUMNS = {
    "_group": "Группа",
    "id": "ID",
    "complex": "Комплект",
    "sheet": "Лист",
    "comment": "Замечание",
    "status": "Статус замечания",
    "doc_status": "Статус листа",
    "author": "Автор",
    "created": "Создано",
    "fix_date": "Наш ответ",
}


def _build_ca_view(df: pd.DataFrame,
                    collapse: bool = True,
                    limit: int | None = 500) -> pd.DataFrame:
    """Собирает таблицу «К листам A и аннулированным».

    collapse=True — скрывает повторы Комплект+Лист (для показа).
    limit — обрезка строк (None — без обрезки, для выгрузки).
    """
    if df.empty:
        return pd.DataFrame()

    view = df

    sort_cols = [c for c in ("_group", "complex", "sheet") if c in view.columns]
    if sort_cols:
        view = view.sort_values(sort_cols, ascending=[True] * len(sort_cols))

    if collapse:
        view = _collapse_repeats(view, ["complex", "sheet"])
    if limit:
        view = view.head(limit)

    cols = [c for c in _CA_COLUMNS if c in view.columns]
    return view[cols].copy().rename(columns=_CA_COLUMNS)


@st.cache_data(ttl=3600, show_spinner=False, max_entries=4)
def _export_rows_to_excel(df: pd.DataFrame) -> bytes:
    """Формирует XLSX из готовой таблицы: заголовки, границы, автофильтр."""
    import numbers

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def _cell_value(v):
        if v is None or pd.isna(v):
            return ""
        if isinstance(v, numbers.Number) and not isinstance(v, complex):
            return v
        if isinstance(v, str):
            return v
        return str(v)

    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Замечания"

    headers = [str(c) for c in df.columns]
    ws.append(headers)

    for row in df.itertuples(index=False, name=None):
        ws.append([_cell_value(v) for v in row])

    width_map = {
        "Группа": 32,
        "ID": 10,
        "Комплект": 24,
        "Лист": 22,
        "Замечание": 70,
        "Статус замечания": 18,
        "Статус листа": 14,
        "Автор": 26,
        "Создано": 18,
        "Наш ответ": 14,
    }
    for c, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width_map.get(name, 16)

    header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0",
                              fill_type="solid")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrapText=True)
        cell.border = border

    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = border

    if "Замечание" in headers:
        col_idx = headers.index("Замечание") + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col_idx).alignment = Alignment(
                wrapText=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    )

    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _render_customer_a_annulled():
    st.markdown("### 📄 Замечания к листам A и аннулированным")
    st.caption(
        "Листы A — утверждены, аннулированные — сняты заказчиком. "
        "Замечания фактически не актуальны, но формально не закрыты. "
        "Ожидаем снятия от заказчика."
    )

    disc_options = _load_discipline_options()
    with st.expander("🎛 Фильтры", expanded=False):
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            sel_disc_labels = st.multiselect(
                "Дисциплина", options=list(disc_options.values()),
                placeholder="Все дисциплины", key="ca_disc",
            )
            sel_disc = [c for c, l in disc_options.items()
                         if l in sel_disc_labels]
        section_options = _load_section_options(
            tuple(sel_disc) if sel_disc else ())
        with c2:
            if section_options:
                sel_section_labels = st.multiselect(
                    "Раздел", options=list(section_options.values()),
                    placeholder="Все разделы", key="ca_section",
                )
                sel_section = [c for c, l in section_options.items()
                                if l in sel_section_labels]
            else:
                sel_section = []
                st.multiselect("Раздел", options=[],
                                placeholder="Не применимы",
                                disabled=True, key="ca_section_empty")
        with c3:
            sel_blocks = st.multiselect(
                "Расположение", options=list(BLOCK_OPTIONS.keys()),
                placeholder="Все", key="ca_blocks",
            )
            block_codes = tuple(BLOCK_OPTIONS[b] for b in sel_blocks)
        kit_options = _load_kit_options(
            tuple(sel_disc) if sel_disc else (),
            tuple(sel_section) if sel_section else (),
            block_codes)
        with c4:
            sel_kit_labels = st.multiselect(
                "Комплект", options=list(kit_options.values()),
                placeholder="Все комплекты", key="ca_kit",
            )
            sel_kit = [c for c, l in kit_options.items()
                        if l in sel_kit_labels]

    df = _load_customer_a_annulled(
        disciplines=tuple(sel_disc) if sel_disc else None,
        kits=tuple(sel_kit) if sel_kit else None,
        sections=tuple(sel_section) if sel_section else None,
        block_codes=block_codes if block_codes else None,
    )

    if df.empty:
        st.success("🎉 Нет замечаний к листам A или аннулированным.")
        return

    # Группа
    df = df.copy()
    df["_group"] = df["holder"].map({
        "customer_a": "🟢 К листам A (утверждён)",
        "annulled":   "⚫ К аннулированным",
    })

    n_a = int((df["holder"] == "customer_a").sum())
    n_annul = int((df["holder"] == "annulled").sum())
    n_total = n_a + n_annul

    st.markdown("##### 📊 Объём")

    c1, c2, c3 = st.columns(3)
    c1.metric(
        "🟢 К листам A (утверждён)", f"{n_a:,}".replace(",", " "),
        help="Лист утверждён — заказчик должен снять замечания.",
    )
    c2.metric(
        "⚫ К аннулированным", f"{n_annul:,}".replace(",", " "),
        help="Лист аннулирован — заказчик должен снять замечания.",
    )
    c3.metric(
        "📊 Итого", f"{n_total:,}".replace(",", " "),
        help="Суммарно замечаний к листам A и аннулированным.",
    )

    st.divider()

    # Фильтр по группе
    st.markdown(f"##### 📋 Список ({len(df):,})".replace(",", " "))

    c1, c2, c3 = st.columns(3)
    with c1:
        sel_group = st.multiselect(
            "Группа",
            options=["🟢 К листам A (утверждён)",
                     "⚫ К аннулированным"],
            placeholder="Все", key="ca_group",
        )
    with c2:
        sel_status = st.multiselect(
            "Статус замечания",
            options=["Новое", "Принято в работу",
                     "Не принято", "К обсуждению", "Выполнено"],
            placeholder="Все", key="ca_status",
        )
    with c3:
        sel_sheet_status = st.multiselect(
            "Статус листа",
            options=["A", "B", "C", "И", "АННУЛИРОВАНО"],
            placeholder="Все", key="ca_sheet_status",
        )

    filtered = df.copy()
    if sel_group:
        filtered = filtered[filtered["_group"].isin(sel_group)]
    if sel_status:
        filtered = filtered[filtered["status"].isin(sel_status)]
    if sel_sheet_status:
        ds_upper = filtered["doc_status"].astype(str).str.strip().str.upper()
        filtered = filtered[ds_upper.isin(sel_sheet_status)]

    if filtered.empty:
        st.warning("По фильтрам нет замечаний.")
        return

    # Таблица для показа: свёртка повторов + обрезка до 500 строк
    table = _build_ca_view(filtered)

    st.caption(
        f"Показано: **{len(table):,}** из **{len(filtered):,}**."
        .replace(",", " ")
    )

    # Выгрузка: полный отфильтрованный список, без свёртки и обрезки
    st.download_button(
        "📥 Выгрузить в Excel",
        data=_export_rows_to_excel(
            _build_ca_view(filtered, collapse=False, limit=None)
        ),
        file_name=(
            f"Листы_A_и_аннулированные_{datetime.now():%Y%m%d_%H%M}.xlsx"
        ),
        mime="application/vnd.openxmlformats-officedocument."
             "spreadsheetml.sheet",
        use_container_width=True,
        key="ca_export",
    )

    st.dataframe(
        table,
        use_container_width=True, hide_index=True, height=600,
        column_config={
            "Группа": st.column_config.TextColumn(width="medium"),
            "ID": st.column_config.NumberColumn(width="small"),
            "Комплект": st.column_config.TextColumn(width="medium"),
            "Лист": st.column_config.TextColumn(width="medium"),
            "Замечание": st.column_config.TextColumn(width="large"),
            "Статус замечания": st.column_config.TextColumn(width="small"),
            "Статус листа": st.column_config.TextColumn(width="small"),
            "Автор": st.column_config.TextColumn(width="medium"),
            "Создано": st.column_config.TextColumn(width="small"),
            "Наш ответ": st.column_config.TextColumn(width="small"),
        },
    )


# ---------------------------------------------------------------------------
#  Точка входа
# ---------------------------------------------------------------------------
def _render_analytics_tab():
    disc_options = _load_discipline_options()

    c1, c2, c3 = st.columns(3)
    with c1:
        sel_disc_labels = st.multiselect(
            "Дисциплина", options=list(disc_options.values()),
            placeholder="Все дисциплины", key="authors_disc",
        )
        sel_disc = [c for c, l in disc_options.items()
                    if l in sel_disc_labels]

    section_options = _load_section_options(
        tuple(sel_disc) if sel_disc else ())

    with c2:
        if section_options:
            sel_section_labels = st.multiselect(
                "Раздел", options=list(section_options.values()),
                placeholder="Все разделы", key="authors_section",
            )
            sel_section = [c for c, l in section_options.items()
                            if l in sel_section_labels]
        else:
            sel_section = []
            st.multiselect("Раздел", options=[],
                            placeholder="Разделы не применимы",
                            disabled=True, key="authors_section_empty")

    kit_options = _load_kit_options(
        tuple(sel_disc) if sel_disc else (),
        tuple(sel_section) if sel_section else ())

    with c3:
        sel_kit_labels = st.multiselect(
            "Комплект", options=list(kit_options.values()),
            placeholder="Все комплекты", key="authors_kit",
        )
        sel_kit = [c for c, l in kit_options.items()
                    if l in sel_kit_labels]

    # Второй ряд: Марка + Статус замечания + Статус листа
    st.write("")
    c4, c5, c6 = st.columns(3)
    with c4:
        sel_marka = st.text_input(
            "Марка комплекта",
            placeholder="например: КЖ1",
            key="authors_marka",
            help="Точное совпадение 2-го блока шифра (марки). "
                 "КЖ1 не найдёт КЖ11 или КЖ10.",
        )
    with c5:
        sel_remark_statuses = st.multiselect(
            "Статус замечания",
            options=REMARK_STATUS_OPTIONS,
            placeholder="Все статусы",
            key="authors_remark_status",
            help="Выполнено, Закрыто, Аннулировано не участвуют в аналитике авторов.",
        )
    with c6:
        sel_sheet_groups = st.multiselect(
            "Статус листа",
            options=SHEET_GROUP_OPTIONS,
            placeholder="Все статусы листов",
            key="authors_sheet_group",
        )

    disc_t = tuple(sel_disc) if sel_disc else ()
    sect_t = tuple(sel_section) if sel_section else ()
    kit_t = tuple(sel_kit) if sel_kit else ()
    sheet_groups_t = tuple(sel_sheet_groups) if sel_sheet_groups else ()
    remark_statuses_t = tuple(sel_remark_statuses) if sel_remark_statuses else ()
    marka_t = (sel_marka or "").strip()

    df_all = _load_all_categorized()
    df_all = _filter_by_hierarchy(df_all, disc_t, sect_t, kit_t,
                                   sheet_groups_t, remark_statuses_t,
                                   marka_t)
    df_ours = df_all[df_all["holder"] == "ours"]

    _render_kpi(df_all, df_ours)

    df = _load_authors(disc_t, sect_t, kit_t, sheet_groups_t,
                        remark_statuses_t, marka_t)
    if df.empty:
        st.info("Нет данных по авторам с учётом фильтров.")
        return

    st.divider()
    _render_authors_table(df)
    st.divider()
    _render_drilldown(df, df_ours)


def render():
    st.header("👤 Авторы замечаний")
    st.caption(
        "Аналитика по инженерам заказчика: сколько выдают, "
        "как быстро рассматривают наши ответы."
    )

    tab_analytics, tab_waiting, tab_a_annulled = st.tabs([
        "📊 Аналитика авторов",
        "🔵 Ждут заказчика",
        "📄 К листам A и аннулированным",
    ])

    with tab_analytics:
        _render_analytics_tab()
    with tab_waiting:
        _render_waiting_review()
    with tab_a_annulled:
        _render_customer_a_annulled()
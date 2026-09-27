# vitro/xlsx_builder.py
"""
Генератор форматированных XLSX-отчётов.
Адаптация старого xlsx_builder.py под текущую БД.

Создаёт:
  - Отчет_по_замечаниям_<дисциплина>.xlsx — по каждому комплекту дисциплины.
  - Отчет_по_замечаниям_Сводка.xlsx — 2 листа: по дисциплинам + по комплектам.
"""

from pathlib import Path
from datetime import datetime
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle

from vitro.sqlite_db import get_conn
from vitro.disciplines import discipline_name


# ---------------------------------------------------------------------------
#  Константы
# ---------------------------------------------------------------------------
CAT_1 = "Принято/корректное"
CAT_2 = "Формальное/нет влияния на СМР"
CAT_3 = "Доп.требование/отсутствует в ТЗ"
CAT_4 = "Не принято/нарушение ТНПА"
CATEGORIES = [CAT_1, CAT_2, CAT_3, CAT_4]

CAT_COLORS = {
    CAT_1: "C6EFCE",
    CAT_2: "FFEB9C",
    CAT_3: "BDD7EE",
    CAT_4: "FFC7CE",
}

STATUS_COLORS = {
    "Закрыто":          "2E7D32",
    "Выполнено":        "A5D6A7",
    "Аннулировано":     "D9D9D9",
    "Новое":            "64B5F6",
    "Принято в работу": "FFD54F",
    "Не принято":       "E57373",
    "К обсуждению":     "FFB74D",
}

PCT_RED = PatternFill("solid", fgColor="FFC7CE")
PCT_YELLOW = PatternFill("solid", fgColor="FFEB9C")
PCT_GREEN = PatternFill("solid", fgColor="C6EFCE")
PCT_RED_FONT = Font(bold=True, color="9C0006")
PCT_YELLOW_FONT = Font(bold=True, color="9C6500")
PCT_GREEN_FONT = Font(bold=True, color="006100")

HEADER_FILL = PatternFill("solid", fgColor="0070C0")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FILL = PatternFill("solid", fgColor="D9E1F2")
TITLE_FONT = Font(bold=True, size=12, color="1F4E78")
LINK_FONT = Font(color="0563C1", underline="single")
THIN_BORDER = Border(
    left=Side(style="thin", color="808080"),
    right=Side(style="thin", color="808080"),
    top=Side(style="thin", color="808080"),
    bottom=Side(style="thin", color="808080"),
)
TOTAL_FILL = PatternFill("solid", fgColor="D9D9D9")


def _excel_sheet_name(code):
    forbidden = ':\\/?*[]'
    out = "".join("_" if ch in forbidden else ch for ch in code)
    return out[:31]


def _today():
    return datetime.now().strftime("%d.%m.%Y")


def _pct_cell(cell, value):
    if value is None:
        return
    try:
        v = float(value)
    except (TypeError, ValueError):
        return
    if v <= 33:
        cell.fill = PCT_RED; cell.font = PCT_RED_FONT
    elif v <= 66:
        cell.fill = PCT_YELLOW; cell.font = PCT_YELLOW_FONT
    else:
        cell.fill = PCT_GREEN; cell.font = PCT_GREEN_FONT


# ---------------------------------------------------------------------------
#  Загрузка данных
# ---------------------------------------------------------------------------
def load_all_data():
    with get_conn() as conn:
        documents = [dict(r) for r in conn.execute(
            """SELECT id, leaf, discipline, section, complex, status, name
               FROM documents""")]
        comments = [dict(r) for r in conn.execute(
            """SELECT id, doc_id, comment, status, author, created,
                      fix_date, category, category_user, category_date
               FROM comments""")]
        complexes = [dict(r) for r in conn.execute(
            "SELECT code, name, discipline FROM complexes")]
    return {"documents": documents, "comments": comments, "complexes": complexes}


# ---------------------------------------------------------------------------
#  Расчёт статистики
# ---------------------------------------------------------------------------
def _is_closed(st):
    return (st or "").strip().upper() in ("ЗАКРЫТО", "ВЫПОЛНЕНО")

def _is_annulled(st):
    return (st or "").strip().upper() == "АННУЛИРОВАНО"

def _is_active(st):
    return (st or "").strip().upper() in (
        "НОВОЕ", "ПРИНЯТО В РАБОТУ", "НЕ ПРИНЯТО", "К ОБСУЖДЕНИЮ",
    )

def _cat_key(cat):
    c = (cat or "").strip()
    if c == CAT_1: return "cat_1"
    if c == CAT_2: return "cat_2"
    if c == CAT_3: return "cat_3"
    if c == CAT_4: return "cat_4"
    return "cat_none"


def compute_stats(documents, comments):
    docs_by_complex = defaultdict(list)
    for d in documents:
        if d.get("complex"):
            docs_by_complex[d["complex"]].append(d)

    doc_to_complex = {d["id"]: d["complex"] for d in documents}
    comments_by_complex = defaultdict(list)
    for c in comments:
        cx = doc_to_complex.get(c["doc_id"])
        if cx:
            comments_by_complex[cx].append(c)

    def empty_row():
        return {
            "total": 0, "closed": 0, "annulled": 0, "active": 0,
            "comment_pct": 0.0,
            "docs_total": 0, "doc_a": 0, "doc_b": 0, "doc_c": 0,
            "doc_info": 0, "doc_other": 0, "doc_pct": 0.0,
            "cat_1": 0, "cat_2": 0, "cat_3": 0, "cat_4": 0, "cat_none": 0,
        }

    by_complex = {}
    for cx in set(docs_by_complex) | set(comments_by_complex):
        cms = comments_by_complex.get(cx, [])
        ds = docs_by_complex.get(cx, [])
        row = empty_row()

        row["total"] = len(cms)
        row["closed"] = sum(1 for c in cms if _is_closed(c["status"]))
        row["annulled"] = sum(1 for c in cms if _is_annulled(c["status"]))
        row["active"] = sum(1 for c in cms if _is_active(c["status"]))
        for c in cms:
            row[_cat_key(c.get("category"))] += 1

        row["comment_pct"] = 0.0 if row["total"] == 0 else round(
            (row["closed"] + row["annulled"]) / row["total"] * 100, 1)

        row["docs_total"] = len(ds)
        row["doc_a"] = sum(1 for d in ds if (d["status"] or "").strip().upper() == "A")
        row["doc_b"] = sum(1 for d in ds if (d["status"] or "").strip().upper() == "B")
        row["doc_c"] = sum(1 for d in ds if (d["status"] or "").strip().upper() == "C")
        row["doc_info"] = sum(1 for d in ds if (d["status"] or "").strip() == "И")
        row["doc_other"] = (row["docs_total"] - row["doc_a"] - row["doc_b"]
                            - row["doc_c"] - row["doc_info"])
        row["doc_pct"] = 0.0 if row["docs_total"] == 0 else round(
            (row["doc_a"] + row["doc_b"] + row["doc_info"]) / row["docs_total"] * 100, 1)

        by_complex[cx] = row

    complex_to_disc = {}
    for d in documents:
        if d.get("complex") and d.get("discipline"):
            complex_to_disc[d["complex"]] = d["discipline"]

    by_discipline = defaultdict(empty_row)
    for cx, st in by_complex.items():
        disc = complex_to_disc.get(cx)
        if not disc:
            continue
        row = by_discipline[disc]
        row["complexes"] = row.get("complexes", 0) + 1
        for k in ("total", "closed", "annulled", "active",
                  "docs_total", "doc_a", "doc_b", "doc_c", "doc_info",
                  "doc_other", "cat_1", "cat_2", "cat_3", "cat_4", "cat_none"):
            row[k] += st.get(k, 0)

    for disc, row in by_discipline.items():
        row["comment_pct"] = 0.0 if row["total"] == 0 else round(
            (row["closed"] + row["annulled"]) / row["total"] * 100, 1)
        row["doc_pct"] = 0.0 if row["docs_total"] == 0 else round(
            (row["doc_a"] + row["doc_b"] + row["doc_info"])
            / row["docs_total"] * 100, 1)

    return {"by_complex": by_complex, "by_discipline": dict(by_discipline)}


# ---------------------------------------------------------------------------
#  Сводка по дисциплине
# ---------------------------------------------------------------------------
def _write_discipline_summary(ws, discipline, complexes, stats, summary_sheet_name):
    ws.merge_cells("A1:U1")
    ws["A1"] = (f"Сводка по дисциплине {discipline} — {discipline_name(discipline)}. "
                f"Обновлено: {_today()}")
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "Комплект", "Наименование",
        "ВСЕГО", "ОСТАЛОСЬ", "ЗАКРЫТО", "АННУЛИРОВАНО",
        "% ВЫПОЛНЕНИЯ (замечания)",
        "Всего листов", "A", "B", "C", "И", "Другой",
        "% ГОТОВНОСТИ (листы)",
        "Принято/корректное", "Формальное", "Доп.требование", "Не принято",
        "Без категории",
    ]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN; c.border = THIN_BORDER
    ws.row_dimensions[2].height = 52

    row = 3
    totals = defaultdict(int)

    for cx in complexes:
        code = cx["code"]
        st = stats["by_complex"].get(code, {})

        cell = ws.cell(row=row, column=1, value=code)
        cell.hyperlink = f"#'{_excel_sheet_name(code)}'!A1"
        cell.font = LINK_FONT; cell.border = THIN_BORDER

        c2 = ws.cell(row=row, column=2, value=cx.get("name", ""))
        c2.border = THIN_BORDER
        c2.alignment = Alignment(wrap_text=True, vertical="top")

        ws.cell(row=row, column=3, value=st.get("total", 0)).border = THIN_BORDER
        ws.cell(row=row, column=4, value=st.get("active", 0)).border = THIN_BORDER
        ws.cell(row=row, column=5, value=st.get("closed", 0)).border = THIN_BORDER
        ws.cell(row=row, column=6, value=st.get("annulled", 0)).border = THIN_BORDER

        c7 = ws.cell(row=row, column=7, value=st.get("comment_pct", 0.0))
        c7.border = THIN_BORDER; _pct_cell(c7, st.get("comment_pct", 0.0))

        ws.cell(row=row, column=8, value=st.get("docs_total", 0)).border = THIN_BORDER
        ws.cell(row=row, column=9, value=st.get("doc_a", 0)).border = THIN_BORDER
        ws.cell(row=row, column=10, value=st.get("doc_b", 0)).border = THIN_BORDER
        ws.cell(row=row, column=11, value=st.get("doc_c", 0)).border = THIN_BORDER
        ws.cell(row=row, column=12, value=st.get("doc_info", 0)).border = THIN_BORDER
        ws.cell(row=row, column=13, value=st.get("doc_other", 0)).border = THIN_BORDER

        c14 = ws.cell(row=row, column=14, value=st.get("doc_pct", 0.0))
        c14.border = THIN_BORDER; _pct_cell(c14, st.get("doc_pct", 0.0))

        ws.cell(row=row, column=15, value=st.get("cat_1", 0)).border = THIN_BORDER
        ws.cell(row=row, column=16, value=st.get("cat_2", 0)).border = THIN_BORDER
        ws.cell(row=row, column=17, value=st.get("cat_3", 0)).border = THIN_BORDER
        ws.cell(row=row, column=18, value=st.get("cat_4", 0)).border = THIN_BORDER
        ws.cell(row=row, column=19, value=st.get("cat_none", 0)).border = THIN_BORDER

        for k in ("total", "active", "closed", "annulled",
                  "docs_total", "doc_a", "doc_b", "doc_c", "doc_info", "doc_other",
                  "cat_1", "cat_2", "cat_3", "cat_4", "cat_none"):
            totals[k] += st.get(k, 0)
        row += 1

    # Итоги
    tr = row
    c = ws.cell(row=tr, column=1, value="ИТОГО")
    c.font = Font(bold=True); c.fill = TOTAL_FILL; c.border = THIN_BORDER
    c = ws.cell(row=tr, column=2, value="Все комплекты дисциплины")
    c.font = Font(bold=True); c.fill = TOTAL_FILL; c.border = THIN_BORDER

    for col, key in [(3, "total"), (4, "active"), (5, "closed"), (6, "annulled"),
                     (8, "docs_total"), (9, "doc_a"), (10, "doc_b"),
                     (11, "doc_c"), (12, "doc_info"), (13, "doc_other"),
                     (15, "cat_1"), (16, "cat_2"), (17, "cat_3"),
                     (18, "cat_4"), (19, "cat_none")]:
        cell = ws.cell(row=tr, column=col, value=totals[key])
        cell.font = Font(bold=True); cell.fill = TOTAL_FILL; cell.border = THIN_BORDER

    if totals["total"] > 0:
        pct_all = round((totals["closed"] + totals["annulled"]) / totals["total"] * 100, 1)
        c7 = ws.cell(row=tr, column=7, value=pct_all)
        _pct_cell(c7, pct_all); c7.border = THIN_BORDER

    if totals["docs_total"] > 0:
        pct_doc = round(
            (totals["doc_a"] + totals["doc_b"] + totals["doc_info"])
            / totals["docs_total"] * 100, 1)
        c14 = ws.cell(row=tr, column=14, value=pct_doc)
        _pct_cell(c14, pct_doc); c14.border = THIN_BORDER

    widths = [24, 50, 8, 10, 10, 14, 16, 10, 8, 8, 8, 8, 8, 16,
              22, 20, 20, 20, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:S{tr - 1}"


# ---------------------------------------------------------------------------
#  Лист комплекта
# ---------------------------------------------------------------------------
def _write_complex_sheet(ws, complex_code, complex_name, docs, comments,
                         summary_sheet_name):
    ws.merge_cells("A1:J1")
    cell = ws["A1"]
    cell.value = "← К сводке"
    cell.hyperlink = f"#'{summary_sheet_name}'!A1"
    cell.font = LINK_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    headers = ["Шифр листа", "Название листа", "ИД", "Замечание",
               "Дата", "Статус", "Автор", "Категория",
               "Кто изменил", "Когда изменил"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN; c.border = THIN_BORDER
    ws.row_dimensions[2].height = 28

    doc_by_id = {d["id"]: d for d in docs}
    comments_sorted = sorted(
        comments,
        key=lambda c: ((doc_by_id.get(c["doc_id"]) or {}).get("leaf", ""), c["id"]),
    )

    row = 3
    prev_leaf = None
    prev_leaf_name = None
    for c in comments_sorted:
        d = doc_by_id.get(c["doc_id"]) or {}
        leaf_short = (d.get("leaf") or "").replace(".pdf", "").replace(".PDF", "")
        leaf_name = d.get("name") or ""

        if leaf_short != prev_leaf:
            ws.cell(row=row, column=1, value=leaf_short); prev_leaf = leaf_short
        if leaf_name != prev_leaf_name:
            ws.cell(row=row, column=2, value=leaf_name); prev_leaf_name = leaf_name

        ws.cell(row=row, column=1).border = THIN_BORDER
        ws.cell(row=row, column=1).alignment = Alignment(vertical="top")
        ws.cell(row=row, column=2).border = THIN_BORDER
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

        ws.cell(row=row, column=3, value=c["id"]).border = THIN_BORDER

        cell4 = ws.cell(row=row, column=4, value=c.get("comment") or "")
        cell4.border = THIN_BORDER
        cell4.alignment = Alignment(wrap_text=True, vertical="top")

        ws.cell(row=row, column=5, value=c.get("created") or "").border = THIN_BORDER

        st = (c.get("status") or "").strip()
        c6 = ws.cell(row=row, column=6, value=st)
        c6.border = THIN_BORDER
        if st in STATUS_COLORS:
            c6.fill = PatternFill("solid", fgColor=STATUS_COLORS[st])

        c7 = ws.cell(row=row, column=7, value=c.get("author") or "")
        c7.border = THIN_BORDER
        c7.alignment = Alignment(wrap_text=True, vertical="top")

        cat = c.get("category") or ""
        c8 = ws.cell(row=row, column=8, value=cat)
        c8.border = THIN_BORDER
        if cat in CAT_COLORS:
            c8.fill = PatternFill("solid", fgColor=CAT_COLORS[cat])

        ws.cell(row=row, column=9, value=c.get("category_user") or "").border = THIN_BORDER
        ws.cell(row=row, column=10, value=c.get("category_date") or "").border = THIN_BORDER
        row += 1

    last_row = row - 1

    widths = [26, 45, 8, 65, 22, 18, 22, 32, 15, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "C3"
    if last_row >= 2:
        ws.auto_filter.ref = f"A2:J{last_row}"

    if last_row >= 3:
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(CATEGORIES) + '"',
            allow_blank=True, showDropDown=False,
        )
        dv.error = "Выберите одну из 4 категорий"
        dv.errorTitle = "Недопустимое значение"
        dv.prompt = "Выберите категорию из списка"
        dv.promptTitle = "Категория замечания"
        ws.add_data_validation(dv)
        dv.add(f"H3:H{last_row}")

        for cat, color in CAT_COLORS.items():
            rule = Rule(type="cellIs", operator="equal",
                        formula=[f'"{cat}"'],
                        dxf=DifferentialStyle(fill=PatternFill(
                            start_color=color, end_color=color,
                            fill_type="solid")))
            ws.conditional_formatting.add(f"H3:H{last_row}", rule)

        for st_val, color in STATUS_COLORS.items():
            rule = Rule(type="cellIs", operator="equal",
                        formula=[f'"{st_val}"'],
                        dxf=DifferentialStyle(fill=PatternFill(
                            start_color=color, end_color=color,
                            fill_type="solid")))
            ws.conditional_formatting.add(f"F3:F{last_row}", rule)


# ---------------------------------------------------------------------------
#  Главная функция
# ---------------------------------------------------------------------------
def build_all_files(output_dir="Отчёты") -> list[Path]:
    """Генерирует все XLSX-отчёты. Возвращает список путей."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    data = load_all_data()
    stats = compute_stats(data["documents"], data["comments"])

    complexes_by_disc = defaultdict(list)
    for cx in data["complexes"]:
        complexes_by_disc[cx["discipline"]].append(cx)

    doc_to_complex = {d["id"]: d["complex"] for d in data["documents"]}
    docs_by_complex = defaultdict(list)
    for d in data["documents"]:
        if d.get("complex"):
            docs_by_complex[d["complex"]].append(d)

    comments_by_complex = defaultdict(list)
    for c in data["comments"]:
        cx = doc_to_complex.get(c["doc_id"])
        if cx:
            comments_by_complex[cx].append(c)

    created = []

    for discipline in sorted(complexes_by_disc):
        complexes = sorted(complexes_by_disc[discipline], key=lambda x: x["code"])
        file_name = f"Отчет_по_замечаниям_{discipline}.xlsx"
        output_path = output_dir / file_name

        wb = Workbook()
        ws_summary = wb.active
        summary_sheet_name = _excel_sheet_name(f"Сводка_{discipline}")
        ws_summary.title = summary_sheet_name

        _write_discipline_summary(ws_summary, discipline, complexes,
                                  stats, summary_sheet_name)

        for cx in complexes:
            code = cx["code"]
            ws = wb.create_sheet(title=_excel_sheet_name(code))
            _write_complex_sheet(ws, code, cx.get("name", ""),
                                 docs_by_complex.get(code, []),
                                 comments_by_complex.get(code, []),
                                 summary_sheet_name)

        wb.save(output_path)
        created.append(output_path)

    # Общий сводный файл
    wb = Workbook()

    # Лист 1: по дисциплинам
    ws = wb.active
    ws.title = "Сводка_по_дисциплинам"
    headers = [
        "Дисциплина", "Наименование", "Комплектов",
        "ВСЕГО", "ОСТАЛОСЬ", "ЗАКРЫТО", "АННУЛИРОВАНО",
        "% ВЫПОЛНЕНИЯ (замечания)",
        "Всего листов", "A", "B", "C", "И", "Другой",
        "% ГОТОВНОСТИ (листы)",
        "Принято", "Формальное", "Доп.треб.", "Не принято", "Без категории",
    ]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN; c.border = THIN_BORDER
    ws.row_dimensions[1].height = 52

    row = 2
    totals_d = defaultdict(int)
    for disc in sorted(stats["by_discipline"]):
        st = stats["by_discipline"][disc]
        ws.cell(row=row, column=1, value=disc).border = THIN_BORDER
        ws.cell(row=row, column=2, value=discipline_name(disc)).border = THIN_BORDER
        ws.cell(row=row, column=3, value=st.get("complexes", 0)).border = THIN_BORDER
        ws.cell(row=row, column=4, value=st.get("total", 0)).border = THIN_BORDER
        ws.cell(row=row, column=5, value=st.get("active", 0)).border = THIN_BORDER
        ws.cell(row=row, column=6, value=st.get("closed", 0)).border = THIN_BORDER
        ws.cell(row=row, column=7, value=st.get("annulled", 0)).border = THIN_BORDER

        c8 = ws.cell(row=row, column=8, value=st.get("comment_pct", 0.0))
        c8.border = THIN_BORDER; _pct_cell(c8, st.get("comment_pct", 0.0))

        ws.cell(row=row, column=9, value=st.get("docs_total", 0)).border = THIN_BORDER
        ws.cell(row=row, column=10, value=st.get("doc_a", 0)).border = THIN_BORDER
        ws.cell(row=row, column=11, value=st.get("doc_b", 0)).border = THIN_BORDER
        ws.cell(row=row, column=12, value=st.get("doc_c", 0)).border = THIN_BORDER
        ws.cell(row=row, column=13, value=st.get("doc_info", 0)).border = THIN_BORDER
        ws.cell(row=row, column=14, value=st.get("doc_other", 0)).border = THIN_BORDER

        c15 = ws.cell(row=row, column=15, value=st.get("doc_pct", 0.0))
        c15.border = THIN_BORDER; _pct_cell(c15, st.get("doc_pct", 0.0))

        ws.cell(row=row, column=16, value=st.get("cat_1", 0)).border = THIN_BORDER
        ws.cell(row=row, column=17, value=st.get("cat_2", 0)).border = THIN_BORDER
        ws.cell(row=row, column=18, value=st.get("cat_3", 0)).border = THIN_BORDER
        ws.cell(row=row, column=19, value=st.get("cat_4", 0)).border = THIN_BORDER
        ws.cell(row=row, column=20, value=st.get("cat_none", 0)).border = THIN_BORDER

        for k in ("complexes", "total", "active", "closed", "annulled",
                  "docs_total", "doc_a", "doc_b", "doc_c", "doc_info", "doc_other",
                  "cat_1", "cat_2", "cat_3", "cat_4", "cat_none"):
            totals_d[k] += st.get(k, 0)
        row += 1

    # Итого
    tr = row
    c = ws.cell(row=tr, column=1, value="ИТОГО")
    c.font = Font(bold=True); c.fill = TOTAL_FILL; c.border = THIN_BORDER
    c = ws.cell(row=tr, column=2, value="Все дисциплины")
    c.font = Font(bold=True); c.fill = TOTAL_FILL; c.border = THIN_BORDER

    for col, key in [(3, "complexes"), (4, "total"), (5, "active"),
                     (6, "closed"), (7, "annulled"),
                     (9, "docs_total"), (10, "doc_a"), (11, "doc_b"),
                     (12, "doc_c"), (13, "doc_info"), (14, "doc_other"),
                     (16, "cat_1"), (17, "cat_2"), (18, "cat_3"),
                     (19, "cat_4"), (20, "cat_none")]:
        cell = ws.cell(row=tr, column=col, value=totals_d[key])
        cell.font = Font(bold=True); cell.fill = TOTAL_FILL; cell.border = THIN_BORDER

    if totals_d["total"] > 0:
        pct = round((totals_d["closed"] + totals_d["annulled"])
                    / totals_d["total"] * 100, 1)
        c8 = ws.cell(row=tr, column=8, value=pct)
        c8.border = THIN_BORDER; _pct_cell(c8, pct)

    if totals_d["docs_total"] > 0:
        pct_d = round(
            (totals_d["doc_a"] + totals_d["doc_b"] + totals_d["doc_info"])
            / totals_d["docs_total"] * 100, 1)
        c15 = ws.cell(row=tr, column=15, value=pct_d)
        c15.border = THIN_BORDER; _pct_cell(c15, pct_d)

    widths = [12, 40, 12, 10, 10, 10, 14, 16, 10, 8, 8, 8, 8, 8, 16,
              14, 14, 14, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:T{tr - 1}"

    # Лист 2: по комплектам
    ws2 = wb.create_sheet(title="Сводка_по_комплектам")
    headers2 = [
        "Дисциплина", "Комплект", "Наименование",
        "ВСЕГО", "ОСТАЛОСЬ", "ЗАКРЫТО", "АННУЛИРОВАНО",
        "% ВЫПОЛНЕНИЯ (замечания)",
        "Всего листов", "A", "B", "C", "И", "Другой",
        "% ГОТОВНОСТИ (листы)",
        "Принято", "Формальное", "Доп.треб.", "Не принято", "Без категории",
    ]
    for i, h in enumerate(headers2, start=1):
        c = ws2.cell(row=1, column=i, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN; c.border = THIN_BORDER
    ws2.row_dimensions[1].height = 52

    items = [(cx["discipline"], cx["code"], cx) for cx in data["complexes"]]
    items.sort(key=lambda t: (t[0], t[1]))

    row = 2
    for disc, code, cx in items:
        st = stats["by_complex"].get(code, {})
        ws2.cell(row=row, column=1, value=disc).border = THIN_BORDER
        ws2.cell(row=row, column=2, value=code).border = THIN_BORDER

        cell3 = ws2.cell(row=row, column=3, value=cx.get("name", ""))
        cell3.border = THIN_BORDER
        cell3.alignment = Alignment(wrap_text=True, vertical="top")

        ws2.cell(row=row, column=4, value=st.get("total", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=5, value=st.get("active", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=6, value=st.get("closed", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=7, value=st.get("annulled", 0)).border = THIN_BORDER

        c8 = ws2.cell(row=row, column=8, value=st.get("comment_pct", 0.0))
        c8.border = THIN_BORDER; _pct_cell(c8, st.get("comment_pct", 0.0))

        ws2.cell(row=row, column=9, value=st.get("docs_total", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=10, value=st.get("doc_a", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=11, value=st.get("doc_b", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=12, value=st.get("doc_c", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=13, value=st.get("doc_info", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=14, value=st.get("doc_other", 0)).border = THIN_BORDER

        c15 = ws2.cell(row=row, column=15, value=st.get("doc_pct", 0.0))
        c15.border = THIN_BORDER; _pct_cell(c15, st.get("doc_pct", 0.0))

        ws2.cell(row=row, column=16, value=st.get("cat_1", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=17, value=st.get("cat_2", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=18, value=st.get("cat_3", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=19, value=st.get("cat_4", 0)).border = THIN_BORDER
        ws2.cell(row=row, column=20, value=st.get("cat_none", 0)).border = THIN_BORDER
        row += 1

    tr2 = row
    totals_c = defaultdict(int)
    for cx in data["complexes"]:
        st = stats["by_complex"].get(cx["code"], {})
        for k in ("total", "active", "closed", "annulled",
                  "docs_total", "doc_a", "doc_b", "doc_c", "doc_info", "doc_other",
                  "cat_1", "cat_2", "cat_3", "cat_4", "cat_none"):
            totals_c[k] += st.get(k, 0)

    c = ws2.cell(row=tr2, column=1, value="ИТОГО")
    c.font = Font(bold=True); c.fill = TOTAL_FILL; c.border = THIN_BORDER
    c = ws2.cell(row=tr2, column=2, value="Все комплекты")
    c.font = Font(bold=True); c.fill = TOTAL_FILL; c.border = THIN_BORDER
    c = ws2.cell(row=tr2, column=3, value="")
    c.fill = TOTAL_FILL; c.border = THIN_BORDER

    for col, key in [(4, "total"), (5, "active"), (6, "closed"), (7, "annulled"),
                     (9, "docs_total"), (10, "doc_a"), (11, "doc_b"),
                     (12, "doc_c"), (13, "doc_info"), (14, "doc_other"),
                     (16, "cat_1"), (17, "cat_2"), (18, "cat_3"),
                     (19, "cat_4"), (20, "cat_none")]:
        cell = ws2.cell(row=tr2, column=col, value=totals_c[key])
        cell.font = Font(bold=True); cell.fill = TOTAL_FILL; cell.border = THIN_BORDER

    if totals_c["total"] > 0:
        pct_all = round((totals_c["closed"] + totals_c["annulled"])
                        / totals_c["total"] * 100, 1)
        c8 = ws2.cell(row=tr2, column=8, value=pct_all)
        c8.border = THIN_BORDER; _pct_cell(c8, pct_all)

    if totals_c["docs_total"] > 0:
        pct_doc_all = round(
            (totals_c["doc_a"] + totals_c["doc_b"] + totals_c["doc_info"])
            / totals_c["docs_total"] * 100, 1)
        c15 = ws2.cell(row=tr2, column=15, value=pct_doc_all)
        c15.border = THIN_BORDER; _pct_cell(c15, pct_doc_all)

    widths2 = [12, 24, 50, 10, 10, 10, 14, 16, 10, 8, 8, 8, 8, 8, 16,
               14, 14, 14, 14, 14]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = f"A1:T{tr2 - 1}"

    summary_path = output_dir / "Отчет_по_замечаниям_Сводка.xlsx"
    wb.save(summary_path)
    created.append(summary_path)

    return created